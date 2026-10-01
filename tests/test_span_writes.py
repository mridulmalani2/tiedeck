"""PLAN.md §5.3: a correction writes the figure's own characters, not its run.

Two defects with one cause, both in the write path every derived **Fix it**
goes through. The figure was addressed by the run holding it, and the run was
replaced:

* in prose a sentence is usually one run, so correcting "EBITDA margin of
  25.8% in FY25A." wrote "24.8%" over **the whole sentence**;
* a figure split across runs had no single run to replace, so it was refused
  -- correctly, since writing the first run left "2,10012m" -- and never
  corrected;
* and one PLAN.md did not list: a table cell reading "26.8%*" with its marker
  in the same run lost the marker.

Each test fails on the write path as it was. They build their own decks and
apply the fix the page would offer, then read the deck back.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

from tieout.learn import learn_from_decks
from tieout.model.loader import load_deck
from tieout.rules.base import clear_caches, run_rules
from tieout_fix import apply_fix, plan_fixes

_PNL = [
    ["$ in millions", "FY24A", "FY25A"],
    ["Revenue", "1,624", "1,770"],
    ["EBITDA", "389", "439"],
    ["EBITDA margin", "24.0%", "24.8%"],
]


def _deck(path: Path, *, table=_PNL, runs: list[tuple[str, bool]] | None = None) -> Path:
    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_table(len(table), len(table[0]), Pt(36), Pt(120), Pt(600), Pt(80))
    for r, row in enumerate(table):
        for c, value in enumerate(row):
            frame.table.cell(r, c).text = value
    if runs:
        box = slide.shapes.add_textbox(Pt(36), Pt(300), Pt(880), Pt(30))
        box.name = "Summary"
        paragraph = box.text_frame.paragraphs[0]
        for text, bold in runs:
            run = paragraph.add_run()
            run.text = text
            run.font.bold = bold
    presentation.save(str(path))
    return path


def _fix_all(path: Path, rule_id: str) -> Path:
    clear_caches()
    deck = load_deck(str(path))
    profile = learn_from_decks([deck], "spans").profile
    fixes = list(plan_fixes(run_rules(deck, profile, include=[rule_id]), profile).values())
    assert fixes, "the rule offered no fix"
    out = path.with_name(f"fixed-{path.name}")
    source = path
    for fix in fixes:
        assert apply_fix(source, out, fix).applied
        source = out
    return out


def _summary_runs(path: Path) -> list[tuple[str, bool]]:
    presentation = Presentation(str(path))
    for shape in presentation.slides[0].shapes:
        if shape.name == "Summary":
            return [
                (run.text, bool(run.font.bold))
                for run in shape.text_frame.paragraphs[0].runs
            ]
    raise AssertionError("no summary box")


def test_a_fix_in_prose_keeps_the_sentence(tmp_path) -> None:
    path = _deck(tmp_path / "prose.pptx", runs=[("EBITDA margin of 25.8% in FY25A.", False)])
    fixed = _fix_all(path, "CO-004")
    assert _summary_runs(fixed) == [("EBITDA margin of 24.8% in FY25A.", False)]


def test_a_figure_split_across_runs_is_written_across_them(tmp_path) -> None:
    """The second run is bold. It keeps its formatting and what it held after
    the figure; the replacement takes the first run's formatting."""
    path = _deck(
        tmp_path / "split.pptx",
        runs=[("EBITDA margin of 2", False), ("5.8% in FY25A.", True)],
    )
    fixed = _fix_all(path, "CO-004")
    assert _summary_runs(fixed) == [
        ("EBITDA margin of 24.8%", False),
        (" in FY25A.", True),
    ]


def test_a_fix_in_a_cell_keeps_a_marker_in_the_same_run(tmp_path) -> None:
    table = [*_PNL[:3], ["EBITDA margin", "24.0%", "26.8%*"]]
    fixed = _fix_all(_deck(tmp_path / "marker.pptx", table=table), "CO-004")
    cell = Presentation(str(fixed)).slides[0].shapes[0].table.cell(3, 2)
    assert cell.text == "24.8%*"


def test_a_figure_that_changed_since_the_check_is_refused(tmp_path) -> None:
    """The span is checked before it is written: a deck edited between the
    check and the click must not have characters replaced by position."""
    path = _deck(tmp_path / "stale.pptx", runs=[("EBITDA margin of 25.8% in FY25A.", False)])
    clear_caches()
    deck = load_deck(str(path))
    profile = learn_from_decks([deck], "spans").profile
    (fix,) = plan_fixes(run_rules(deck, profile, include=["CO-004"]), profile).values()

    edited = Presentation(str(path))
    for shape in edited.slides[0].shapes:
        if shape.name == "Summary":
            shape.text_frame.paragraphs[0].runs[0].text = "EBITDA margin of 26.1% in FY25A."
    edited.save(str(path))

    report = apply_fix(path, tmp_path / "out.pptx", fix)
    assert not report.applied
    assert "no longer reads" in report.detail
