"""tieout_ui.canvas: the shape tree the live surface draws instead of a photo.

Each test below is one of PLAN.md section 6's edge cases, reproduced in the
smallest deck that contains it: a shape inside a group, a rotated shape, a
table, a chart, SmartArt, a shape off the canvas or with zero area, and a
slide whose size is not 960x540. The assertion in every case is the same
kind: the surface must describe the shape honestly -- its true geometry, and
a stated reason wherever it cannot draw the content -- never a guess.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest

from tests.test_smartart import _deck_with_smartart
from tieout.model.loader import load_deck
from tieout_ui.canvas import canvas_view


def _presentation(width_pt: float = 960.0, height_pt: float = 540.0):
    from pptx import Presentation
    from pptx.util import Emu

    presentation = Presentation()
    presentation.slide_width = Emu(round(width_pt * 12700))
    presentation.slide_height = Emu(round(height_pt * 12700))
    return presentation


def _blank(presentation):
    return presentation.slides.add_slide(presentation.slide_layouts[6])


@pytest.fixture
def geometry_deck(tmp_path: Path):
    """One slide: a loose shape, a shape inside a group, a rotated shape, a
    shape parked off the canvas, and a shape with zero area."""
    from pptx.util import Emu, Pt

    presentation = _presentation()
    slide = _blank(presentation)

    loose = slide.shapes.add_textbox(Pt(60), Pt(80), Pt(300), Pt(40))
    loose.name = "Loose"
    loose.text_frame.text = "Not in a group."

    group = slide.shapes.add_group_shape()
    group.name = "Group"
    inner = group.shapes.add_textbox(Pt(400), Pt(200), Pt(200), Pt(40))
    inner.name = "Inside"
    inner.text_frame.text = "In a group."
    # add_group_shape() leaves the group's own xfrm at zero extent, which the
    # loader's group-transform math divides by; give it a real box.
    group.left, group.top, group.width, group.height = Pt(400), Pt(200), Pt(200), Pt(40)

    rotated = slide.shapes.add_textbox(Pt(600), Pt(300), Pt(120), Pt(40))
    rotated.name = "Rotated"
    rotated.rotation = 45.0
    rotated.text_frame.text = "Tilted."

    off_canvas = slide.shapes.add_textbox(Pt(-200), Pt(50), Pt(100), Pt(30))
    off_canvas.name = "Off canvas"
    off_canvas.text_frame.text = "Parked outside the slide."

    zero = slide.shapes.add_shape(1, Emu(0), Emu(0), Emu(0), Emu(0))
    zero.name = "Zero area"

    path = tmp_path / "geometry.pptx"
    presentation.save(str(path))
    return path


@pytest.fixture
def table_deck(tmp_path: Path):
    from pptx.util import Pt

    presentation = _presentation()
    slide = _blank(presentation)
    graphic_frame = slide.shapes.add_table(2, 2, Pt(60), Pt(60), Pt(400), Pt(120))
    graphic_frame.name = "A table"
    table = graphic_frame.table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "$100"

    path = tmp_path / "table.pptx"
    presentation.save(str(path))
    return path


@pytest.fixture
def chart_deck(tmp_path: Path):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Pt

    presentation = _presentation()
    slide = _blank(presentation)
    data = CategoryChartData()
    data.categories = ["FY24A", "FY25E"]
    data.add_series("Revenue", (10.0, 12.0))
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(60), Pt(60), Pt(500), Pt(280), data
    )
    frame.name = "A chart"
    frame.chart.has_title = True
    frame.chart.chart_title.text_frame.text = "Revenue growth"

    path = tmp_path / "chart.pptx"
    presentation.save(str(path))
    return path


@pytest.fixture
def smartart_deck(tmp_path: Path) -> Path:
    return _deck_with_smartart(tmp_path / "smartart.pptx", ["Sign", "Diligence", "Close"])


@pytest.fixture
def odd_size_deck(tmp_path: Path):
    """A slide sized 959.976 x 539.986pt -- the client deck's own dimensions,
    not the round 960x540 every fixture elsewhere in the suite uses."""
    from pptx.util import Pt

    presentation = _presentation(width_pt=959.976, height_pt=539.986)
    slide = _blank(presentation)
    box = slide.shapes.add_textbox(Pt(40), Pt(40), Pt(400), Pt(40))
    box.name = "Headline"
    box.text_frame.text = "Odd-sized canvas."

    path = tmp_path / "odd.pptx"
    presentation.save(str(path))
    return path


def _shape(view: dict[str, Any], name: str) -> dict[str, Any]:
    return next(s for s in view["shapes"] if s["name"] == name)


# --------------------------------------------------------------------------------------
# Groups: children already flattened to slide space
# --------------------------------------------------------------------------------------


def test_a_shape_inside_a_group_is_drawn_at_its_true_slide_position(
    geometry_deck, reference_profile
):
    """The model flattens group transforms before the surface ever sees a
    shape, so no group-transform math belongs in the renderer -- the child's
    slide-space box is exactly what the group's own placement put it at."""
    deck = load_deck(geometry_deck)
    view = canvas_view(deck.slides[0], deck, reference_profile)

    inside = _shape(view, "Group > Inside")
    assert inside["left_pt"] == pytest.approx(400.0)
    assert inside["top_pt"] == pytest.approx(200.0)
    assert inside["group_path"] == ["Group"]


def test_a_shape_inside_a_plain_group_is_reported_movable(
    geometry_deck, reference_profile
):
    """tieout_fix writes into a plain group's own coordinate space rather
    than refusing outright, and the flags here have to agree with it -- a
    second copy of "can this be written" that disagreed would offer a button
    that then failed, or refuse one that would have worked."""
    deck = load_deck(geometry_deck)
    view = canvas_view(deck.slides[0], deck, reference_profile)

    inside = _shape(view, "Group > Inside")
    assert inside["editable"]["movable"] is True
    assert inside["editable"]["move_refused"] == ""
    assert inside["editable"]["resizable"] is True
    assert inside["editable"]["resize_refused"] == ""

    loose = _shape(view, "Loose")
    assert loose["editable"]["movable"] is True
    assert loose["editable"]["move_refused"] == ""


def test_a_shape_inside_a_rotated_group_is_refused_for_move_and_resize(
    tmp_path, reference_profile
):
    from pptx.util import Pt

    presentation = _presentation()
    slide = _blank(presentation)
    group = slide.shapes.add_group_shape()
    group.name = "Group"
    inner = group.shapes.add_textbox(Pt(400), Pt(200), Pt(200), Pt(40))
    inner.name = "Inside"
    inner.text_frame.text = "Rotated group."
    group.rotation = 30.0
    path = tmp_path / "rotated-group.pptx"
    presentation.save(str(path))

    deck = load_deck(path)
    view = canvas_view(deck.slides[0], deck, reference_profile)
    inside = _shape(view, "Group > Inside")
    assert inside["editable"]["movable"] is False
    assert "rotat" in inside["editable"]["move_refused"]
    assert inside["editable"]["resizable"] is False
    assert "rotat" in inside["editable"]["resize_refused"]


# --------------------------------------------------------------------------------------
# Rotation
# --------------------------------------------------------------------------------------


def test_a_rotated_shape_carries_its_rotation_and_an_unrotated_box(
    geometry_deck, reference_profile
):
    """The surface hands over the stored box and the rotation separately,
    rather than pre-rotating a bounding box itself -- the frontend needs the
    real box to draw a rotated selection handle, not an axis-aligned photo of
    one."""
    deck = load_deck(geometry_deck)
    view = canvas_view(deck.slides[0], deck, reference_profile)

    rotated = _shape(view, "Rotated")
    assert rotated["rotation"] == pytest.approx(45.0)
    assert rotated["width_pt"] == pytest.approx(120.0)
    assert rotated["height_pt"] == pytest.approx(40.0)
    assert rotated["editable"]["movable"] is True


# --------------------------------------------------------------------------------------
# Off-canvas, and zero area
# --------------------------------------------------------------------------------------


def test_a_shape_parked_off_the_canvas_is_still_reported_at_its_true_position(
    geometry_deck, reference_profile
):
    deck = load_deck(geometry_deck)
    view = canvas_view(deck.slides[0], deck, reference_profile)

    off_canvas = _shape(view, "Off canvas")
    assert off_canvas["left_pt"] == pytest.approx(-200.0)


def test_a_zero_area_shape_does_not_crash_the_serialiser(geometry_deck, reference_profile):
    deck = load_deck(geometry_deck)
    view = canvas_view(deck.slides[0], deck, reference_profile)

    zero = _shape(view, "Zero area")
    assert zero["width_pt"] == 0.0
    assert zero["height_pt"] == 0.0


# --------------------------------------------------------------------------------------
# Tables: drawn for real, but as one frame
# --------------------------------------------------------------------------------------


def test_a_table_is_drawn_as_a_read_only_grid(table_deck, reference_profile):
    deck = load_deck(table_deck)
    view = canvas_view(deck.slides[0], deck, reference_profile)

    node = _shape(view, "A table")
    assert node["table"] is not None
    assert node["table"]["row_count"] == 2
    assert node["table"]["column_count"] == 2
    header = next(
        c for c in node["table"]["cells"] if c["row"] == 0 and c["column"] == 0
    )
    assert header["paragraphs"][0]["runs"][0]["text"] == "Metric"
    # Selection stops at the frame: nothing marks a cell individually editable,
    # because there is no cell-level write-back for the editor to use.
    assert node["editable"]["text_editable"] is False
    assert node["editable"]["movable"] is True


# --------------------------------------------------------------------------------------
# Charts: a placeholder at the true geometry, never a redraw
# --------------------------------------------------------------------------------------


def test_a_chart_is_a_labelled_placeholder_at_its_true_geometry(chart_deck, reference_profile):
    deck = load_deck(chart_deck)
    view = canvas_view(deck.slides[0], deck, reference_profile)

    node = _shape(view, "A chart")
    assert node["kind"] == "chart"
    assert node["placeholder_content"] is not None
    assert "does not model chart layout" in node["placeholder_content"]["reason"]
    assert node["placeholder_content"]["title"] == "Revenue growth"
    # The frame itself is still a real box someone can move.
    assert node["width_pt"] == pytest.approx(500.0)
    assert node["editable"]["movable"] is True
    assert node["editable"]["text_editable"] is False


# --------------------------------------------------------------------------------------
# SmartArt: a stated refusal, never a guess at unmodelled geometry
# --------------------------------------------------------------------------------------


def test_smartart_is_a_labelled_placeholder_naming_what_is_not_modelled(
    smartart_deck, reference_profile
):
    deck = load_deck(smartart_deck)
    # The diagram is on whichever slide _deck_with_smartart put it on; found by
    # kind rather than by name, which python-pptx never set for an injected
    # part.
    view = None
    for candidate in deck.slides:
        candidate_view = canvas_view(candidate, deck, reference_profile)
        if any(s["kind"] == "smartart" for s in candidate_view["shapes"]):
            view = candidate_view
            break
    assert view is not None, "no SmartArt shape was found on any slide"

    node = next(s for s in view["shapes"] if s["kind"] == "smartart")
    assert node["placeholder_content"] is not None
    assert "diagram part" in node["placeholder_content"]["reason"]
    assert node["placeholder_content"]["labels"] == ["Sign", "Diligence", "Close"]
    assert node["editable"]["text_editable"] is False
    # Geometry is still real: the frame can be moved even though its content
    # cannot be drawn or edited.
    assert node["editable"]["movable"] is True


# --------------------------------------------------------------------------------------
# A slide whose size is not 960x540
# --------------------------------------------------------------------------------------


def test_an_odd_sized_slide_reports_its_own_dimensions(odd_size_deck, reference_profile):
    deck = load_deck(odd_size_deck)
    view = canvas_view(deck.slides[0], deck, reference_profile)

    assert view["width_pt"] == pytest.approx(959.976, abs=0.01)
    assert view["height_pt"] == pytest.approx(539.986, abs=0.01)
    headline = _shape(view, "Headline")
    assert headline["left_pt"] == pytest.approx(40.0)


# --------------------------------------------------------------------------------------
# The whole tree is always JSON-serialisable
# --------------------------------------------------------------------------------------


def test_every_fixture_deck_serialises_to_json(
    geometry_deck, table_deck, chart_deck, smartart_deck, odd_size_deck, reference_profile
):
    import json

    for path in (geometry_deck, table_deck, chart_deck, smartart_deck, odd_size_deck):
        deck = load_deck(path)
        for slide in deck.slides:
            json.dumps(canvas_view(slide, deck, reference_profile))
