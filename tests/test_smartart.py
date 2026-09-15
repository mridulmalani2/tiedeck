"""SmartArt: the text inside it, and nothing else.

A ``dgm`` graphic frame carries no text of its own. Every label lives in a
separate diagram part, reachable only through a relationship, which is why a
draft marker or a client's name inside SmartArt was invisible to every text rule
in the tool -- on the one kind of shape a process or structure slide is made of.

What is *not* modelled matters as much. The diagram's internal geometry is laid
out by an algorithm in its layout part, and its colours come from a colour part;
neither is read, so no layout or brand rule is invited to measure a box it
cannot see. The frame's own box still participates in canvas and margin checks,
as it always did.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from tieout.learn import learn_from_decks
from tieout.model.loader import load_deck
from tieout.rules.base import clear_caches, run_rules

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _point(model_id: str, text: str) -> str:
    return (
        f'<dgm:pt modelId="{model_id}"><dgm:prSet/><dgm:spPr/>'
        f'<dgm:t xmlns:a="{_A}"><a:bodyPr/><a:lstStyle/>'
        f'<a:p><a:r><a:rPr lang="en-GB"/><a:t>{text}</a:t></a:r></a:p>'
        f"</dgm:t></dgm:pt>"
    )


def _deck_with_smartart(path: Path, labels: list[str]) -> Path:
    """A one-slide deck carrying a real SmartArt graphic.

    Assembled by injecting the diagram part and its relationship, because
    python-pptx cannot create SmartArt and a fixture that faked the shape would
    not exercise the part resolution this is all about.
    """
    from pptx import Presentation
    from pptx.util import Emu, Pt

    base = path.with_suffix(".base.pptx")
    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Pt(42), Pt(40), Pt(600), Pt(40)).text_frame.text = (
        "Transaction process overview"
    )
    presentation.save(str(base))

    points = "".join(_point(f"{{ID{i}}}", label) for i, label in enumerate(labels, 1))
    # A transition point carries connector formatting and never content. Its
    # empty text body must not read as a blank label.
    points += (
        f'<dgm:pt modelId="{{TRANS}}" type="sibTrans" cxnId="{{C1}}">'
        f'<dgm:prSet/><dgm:spPr/><dgm:t xmlns:a="{_A}"><a:bodyPr/><a:lstStyle/>'
        f'<a:p><a:endParaRPr lang="en-GB"/></a:p></dgm:t></dgm:pt>'
    )
    data = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<dgm:dataModel xmlns:dgm="{_DGM}"><dgm:ptLst>{points}</dgm:ptLst>'
        "<dgm:cxnLst/><dgm:bg/><dgm:whole/></dgm:dataModel>"
    )
    frame = (
        f'<p:graphicFrame xmlns:p="{_P}" xmlns:a="{_A}" xmlns:r="{_R}">'
        '<p:nvGraphicFramePr><p:cNvPr id="9" name="Diagram 1"/>'
        "<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>"
        '<p:xfrm><a:off x="762000" y="1524000"/>'
        '<a:ext cx="7620000" cy="3048000"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{_DGM}">'
        f'<dgm:relIds xmlns:dgm="{_DGM}" r:dm="rId99" r:lo="rId99" '
        'r:qs="rId99" r:cs="rId99"/>'
        "</a:graphicData></a:graphic></p:graphicFrame>"
    )

    with zipfile.ZipFile(base) as source:
        items = {name: source.read(name) for name in source.namelist()}

    items["ppt/slides/slide1.xml"] = (
        items["ppt/slides/slide1.xml"]
        .decode("utf-8")
        .replace("</p:spTree>", frame + "</p:spTree>")
        .encode("utf-8")
    )
    rels = "ppt/slides/_rels/slide1.xml.rels"
    items[rels] = (
        items[rels]
        .decode("utf-8")
        .replace(
            "</Relationships>",
            '<Relationship Id="rId99" Type="'
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
            'diagramData" Target="../diagrams/data1.xml"/></Relationships>',
        )
        .encode("utf-8")
    )
    items["ppt/diagrams/data1.xml"] = data.encode("utf-8")
    items["[Content_Types].xml"] = (
        items["[Content_Types].xml"]
        .decode("utf-8")
        .replace(
            "</Types>",
            '<Override PartName="/ppt/diagrams/data1.xml" ContentType="'
            "application/vnd.openxmlformats-officedocument.drawingml."
            'diagramData+xml"/></Types>',
        )
        .encode("utf-8")
    )

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, payload in items.items():
            out.writestr(name, payload)
    base.unlink()
    return path


@pytest.fixture
def smartart_deck(tmp_path):
    return load_deck(
        _deck_with_smartart(
            tmp_path / "smartart.pptx",
            ["Due diligence", "TBD confirm timing", "Signing"],
        )
    )


def test_the_labels_inside_a_diagram_are_read(smartart_deck):
    shape = next(s for s in smartart_deck.slides[0].all_shapes() if s.kind == "smartart")
    assert shape.diagram_text == ("Due diligence", "TBD confirm timing", "Signing")


def test_a_transition_point_is_not_read_as_a_blank_label(smartart_deck):
    """``parTrans`` and ``sibTrans`` points hold connector formatting. A reader
    sees nothing for them, so an empty string from one would be a label the
    deck does not have."""
    shape = next(s for s in smartart_deck.slides[0].all_shapes() if s.kind == "smartart")
    assert all(label.strip() for label in shape.diagram_text)


def test_diagram_text_is_part_of_the_shape_s_text(smartart_deck):
    shape = next(s for s in smartart_deck.slides[0].all_shapes() if s.kind == "smartart")
    assert "Due diligence" in shape.text
    assert shape.has_text


def test_a_draft_marker_inside_smartart_is_reported(smartart_deck):
    """The point of all of it. A process slide is SmartArt, and "TBD" left in
    one used to go out."""
    profile = learn_from_decks([smartart_deck], "acme").profile
    clear_caches()
    result = run_rules(smartart_deck, profile, include=["HY-001"])

    assert [f.rule_id for f in result.findings] == ["HY-001"]
    assert "TBD" in (result.findings[0].measured or "")


def test_a_diagram_is_never_treated_as_a_text_frame(smartart_deck):
    """Its labels are checkable; its internal geometry is not modelled. Folding
    the text into ``text_frame_paragraphs`` would have invited every layout rule
    to measure boxes that are not there."""
    shape = next(s for s in smartart_deck.slides[0].all_shapes() if s.kind == "smartart")
    assert shape.text_frame_paragraphs == ()


def test_a_deck_with_no_diagram_part_is_unaffected(clean_deck):
    assert all(
        shape.diagram_text == ()
        for slide in clean_deck.slides
        for shape in slide.all_shapes()
    )
