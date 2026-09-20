"""Consistency rule tests.

These rules are the ones a banker would actually thank you for, and also the
ones most able to embarrass the tool: a false "these two figures disagree" on a
deck where they legitimately differ is worse than silence, because it sends
someone to check a number that was right.

So alongside the seeded-defect and clean-deck tests, this module builds small
purpose-built decks to pin the label-matching behaviour exactly — which pairs
compare, which are too generic to compare, and which disagreements are
arithmetic rather than opinion.
"""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

from tests.conftest import assert_silent_on_clean, findings_for, slide_indices
from tieout.figures import build_index
from tieout.fixtures.spec import default_defects
from tieout.model.loader import load_deck
from tieout.rules.base import run_rules
from tieout.rules.consistency import _is_total_label, normalise_label

SEEDED = {d.rule_id: d.slide_index for d in default_defects() if d.variant is None}
CONSISTENCY_RULE_IDS = ("CO-001", "CO-002", "CO-003")


def _run(deck, profile, rule_id):
    return run_rules(deck, profile, include=[rule_id])


# --------------------------------------------------------------------------------------
# Purpose-built decks
# --------------------------------------------------------------------------------------


def _deck_with_tables(path, tables):
    """Build a deck of one-table slides. ``tables`` is a list of row lists."""
    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    for index, rows in enumerate(tables):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_table(
            len(rows), len(rows[0]), Pt(36), Pt(120), Pt(888), Pt(24 * len(rows))
        )
        frame.name = f"Table {index + 1}"
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                frame.table.cell(row_index, column_index).text = str(value)
    presentation.save(str(path))
    return load_deck(path)


_PROJECTIONS = [
    ["Fiscal year", "Revenue", "EBITDA", "Margin"],
    ["2024A", "1,562", "263", "16.8%"],
    ["2025A", "1,908", "351", "18.4%"],
]


@pytest.fixture
def agreeing(tmp_path):
    """Two tables restating the same figures identically. Must be silent."""
    return _deck_with_tables(
        tmp_path / "agree.pptx",
        [
            _PROJECTIONS,
            [["Fiscal year", "EBITDA"], ["2025A", "351"]],
        ],
    )


@pytest.fixture
def disagreeing(tmp_path):
    """Two tables restating the same figure differently."""
    return _deck_with_tables(
        tmp_path / "disagree.pptx",
        [
            _PROJECTIONS,
            [["Fiscal year", "EBITDA"], ["2025A", "375"]],
        ],
    )


# --------------------------------------------------------------------------------------
# Registration and documentation
# --------------------------------------------------------------------------------------


def test_every_consistency_rule_is_registered_and_documented():
    from tieout.rules.base import load_all_rules

    registry = load_all_rules()
    for rule_id in CONSISTENCY_RULE_IDS:
        assert rule_id in registry
        rule = registry[rule_id]
        assert rule.category == "consistency"
        assert rule.summary
        assert rule.__doc__ and "false-positive mode" in rule.__doc__


def test_the_consistency_rules_need_nothing_from_the_profile():
    """Internal arithmetic is not a house style, so these rules run on any deck
    with or without a learned profile."""
    from tieout.rules.base import load_all_rules

    registry = load_all_rules()
    for rule_id in CONSISTENCY_RULE_IDS:
        assert registry[rule_id].requires == ()


# --------------------------------------------------------------------------------------
# CO-001 contradictory figure
# --------------------------------------------------------------------------------------


def test_co001_catches_the_seeded_contradiction(dirty_deck, reference_profile):
    findings = findings_for(_run(dirty_deck, reference_profile, "CO-001"), "CO-001")
    assert SEEDED["CO-001"] in slide_indices(findings)
    message = next(f for f in findings if f.slide_index == SEEDED["CO-001"]).message
    assert "375" in message and "351" in message
    assert "slide 6" in message, "it must say where the other figure lives"


def test_co001_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    assert_silent_on_clean(_run(clean_deck, reference_profile, "CO-001"), "CO-001")


def test_co001_says_nothing_when_two_tables_agree(agreeing, reference_profile):
    assert findings_for(_run(agreeing, reference_profile, "CO-001"), "CO-001") == []


def test_co001_reports_when_two_tables_disagree(disagreeing, reference_profile):
    findings = findings_for(_run(disagreeing, reference_profile, "CO-001"), "CO-001")
    assert len(findings) == 1
    assert findings[0].measured == "375"


def test_co001_ignores_repetition_inside_one_table(tmp_path, reference_profile):
    """A label repeated within a single table is a layout artefact, not two
    statements of the same fact."""
    deck = _deck_with_tables(
        tmp_path / "one.pptx",
        [
            [
                ["Fiscal year", "EBITDA"],
                ["2025A", "351"],
                ["2025A", "375"],
            ]
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-001"), "CO-001") == []


def test_co001_will_not_compare_a_percentage_with_a_multiple(
    tmp_path, reference_profile
):
    """Under one label, 9.2x and 9.2% are different quantities. Comparing them
    would report a contradiction between two numbers that were never the same
    kind of thing."""
    deck = _deck_with_tables(
        tmp_path / "suffix.pptx",
        [
            [["Company", "Multiple"], ["Calderwood", "9.2x"]],
            [["Company", "Multiple"], ["Calderwood", "17.4%"]],
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-001"), "CO-001") == []


def test_co001_will_not_compare_on_a_generic_label(tmp_path, reference_profile):
    """"Value" identifies nothing. Keying on it is how a consistency rule starts
    reporting every table in a deck against every other."""
    deck = _deck_with_tables(
        tmp_path / "generic.pptx",
        [
            [["Item", "Value"], ["Other", "100"]],
            [["Item", "Value"], ["Other", "250"]],
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-001"), "CO-001") == []


def test_co001_leaves_a_pure_scale_error_to_co002(tmp_path, reference_profile):
    """A clean factor of 1000 is a units problem with a different fix, so it is
    reported once by CO-002 rather than twice by both rules."""
    deck = _deck_with_tables(
        tmp_path / "scale.pptx",
        [
            [["Fiscal year", "Revenue"], ["2025A", "1,908"]],
            [["Fiscal year", "Revenue"], ["2025A", "1,908,000"]],
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-001"), "CO-001") == []
    assert findings_for(_run(deck, reference_profile, "CO-002"), "CO-002")


def test_co001_skips_cells_that_are_not_numbers(tmp_path, reference_profile):
    deck = _deck_with_tables(
        tmp_path / "placeholders.pptx",
        [
            [["Objective", "Status"], ["Capital raised", "n.a."]],
            [["Objective", "Status"], ["Capital raised", "Full"]],
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-001"), "CO-001") == []


# --------------------------------------------------------------------------------------
# CO-002 scale mismatch
# --------------------------------------------------------------------------------------


def test_co002_catches_the_seeded_unit_error(dirty_deck, reference_profile):
    findings = findings_for(_run(dirty_deck, reference_profile, "CO-002"), "CO-002")
    assert SEEDED["CO-002"] in slide_indices(findings)
    message = next(f for f in findings if f.slide_index == SEEDED["CO-002"]).message
    assert "wrong unit" in message
    assert "1,000" in message or "factor of 1,000" in message


def test_co002_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    assert_silent_on_clean(_run(clean_deck, reference_profile, "CO-002"), "CO-002")


@pytest.mark.parametrize(
    ("other", "expected"),
    [("1,908,000", True), ("190,800", True), ("1,920", False), ("3,816", False)],
)
def test_co002_only_fires_on_a_clean_factor(
    tmp_path, reference_profile, other, expected
):
    """A thousand-fold or hundred-fold difference is a unit error. A doubling is a
    disagreement, and CO-001's business."""
    deck = _deck_with_tables(
        tmp_path / f"factor-{other.replace(',', '')}.pptx",
        [
            [["Fiscal year", "Revenue"], ["2025A", "1,908"]],
            [["Fiscal year", "Revenue"], ["2025A", other]],
        ],
    )
    fired = bool(findings_for(_run(deck, reference_profile, "CO-002"), "CO-002"))
    assert fired is expected


# --------------------------------------------------------------------------------------
# CO-003 totals
# --------------------------------------------------------------------------------------


def test_co003_catches_the_seeded_broken_total(dirty_deck, reference_profile):
    findings = findings_for(_run(dirty_deck, reference_profile, "CO-003"), "CO-003")
    assert SEEDED["CO-003"] in slide_indices(findings)
    message = next(f for f in findings if f.slide_index == SEEDED["CO-003"]).message
    assert "9,900" in message and "9,828" in message


def test_co003_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    """The clean deck's projections table carries a cumulative row whose
    arithmetic is correct, so this is a real verification rather than a rule
    passing because it found nothing to check."""
    assert_silent_on_clean(_run(clean_deck, reference_profile, "CO-003"), "CO-003")


def test_the_clean_deck_actually_contains_a_total_row_to_verify(clean_deck):
    """Guards the test above from silently becoming vacuous."""
    totals = [
        cell
        for slide in clean_deck.slides
        for shape in slide.tables
        if shape.table
        for cell in shape.table.cells
        if cell.column == 0 and _is_total_label(normalise_label(cell.text))
    ]
    assert totals, "the fixture is supposed to carry a total row for CO-003 to check"


def test_co003_tolerates_rounding(tmp_path, reference_profile):
    """Five figures each rounded to the nearest whole number can legitimately sum
    2.5 away from their rounded total. A rule demanding exactness would report
    every correctly-rounded table in banking."""
    deck = _deck_with_tables(
        tmp_path / "rounding.pptx",
        [
            [
                ["Segment", "Revenue"],
                ["North", "10"],
                ["South", "10"],
                ["East", "10"],
                ["West", "10"],
                ["Total", "42"],
            ]
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-003"), "CO-003") == []


def test_co003_reports_a_total_beyond_the_rounding_tolerance(
    tmp_path, reference_profile
):
    deck = _deck_with_tables(
        tmp_path / "broken.pptx",
        [
            [
                ["Segment", "Revenue"],
                ["North", "10"],
                ["South", "10"],
                ["East", "10"],
                ["West", "10"],
                ["Total", "60"],
            ]
        ],
    )
    findings = findings_for(_run(deck, reference_profile, "CO-003"), "CO-003")
    assert len(findings) == 1
    assert "60" in findings[0].message and "40" in findings[0].message


def test_co003_does_not_add_up_percentages(tmp_path, reference_profile):
    """A "Total" over a margin column is a weighted average, not a sum, and
    adding the percentages would report every margin table ever drawn."""
    deck = _deck_with_tables(
        tmp_path / "percent.pptx",
        [
            [
                ["Segment", "Margin"],
                ["North", "15.0%"],
                ["South", "16.0%"],
                ["East", "17.0%"],
                ["Total", "16.0%"],
            ]
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-003"), "CO-003") == []


def test_co003_needs_at_least_three_addends(tmp_path, reference_profile):
    """Two numbers and a total is as often a subtraction or a comparison."""
    deck = _deck_with_tables(
        tmp_path / "two.pptx",
        [[["Segment", "Revenue"], ["North", "10"], ["Total", "40"]]],
    )
    assert findings_for(_run(deck, reference_profile, "CO-003"), "CO-003") == []


def test_co003_skips_a_column_holding_non_numeric_rows(tmp_path, reference_profile):
    """If a column mixes prose and figures it is not a sum, whatever the total row
    says."""
    deck = _deck_with_tables(
        tmp_path / "mixed.pptx",
        [
            [
                ["Segment", "Revenue"],
                ["North", "10"],
                ["South", "Pending"],
                ["East", "10"],
                ["West", "10"],
                ["Total", "60"],
            ]
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-003"), "CO-003") == []


def test_co003_ignores_a_label_that_merely_starts_with_total(
    tmp_path, reference_profile
):
    """"Total addressable market" is a metric, not an assertion of arithmetic."""
    deck = _deck_with_tables(
        tmp_path / "tam.pptx",
        [
            [
                ["Segment", "Value"],
                ["North", "10"],
                ["South", "10"],
                ["East", "10"],
                ["Total addressable market", "900"],
            ]
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-003"), "CO-003") == []


def test_co003_excludes_a_subtotal_from_the_grand_total(tmp_path, reference_profile):
    """A grand total sums the line items, not the line items plus a subtotal."""
    deck = _deck_with_tables(
        tmp_path / "subtotal.pptx",
        [
            [
                ["Segment", "Revenue"],
                ["North", "10"],
                ["South", "10"],
                ["Sum", "20"],
                ["East", "10"],
                ["West", "10"],
                ["Total", "40"],
            ]
        ],
    )
    assert findings_for(_run(deck, reference_profile, "CO-003"), "CO-003") == []


# --------------------------------------------------------------------------------------
# Shared machinery and the category guard
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("EBITDA(1)", "ebitda"),
        ("EBITDA ", "ebitda"),
        ("Revenue*", "revenue"),
        ("Free cash flow", "free cash flow"),
        ("FY24", "fy24"),
        ("Q1 2026", "q1 2026"),
        ("Top 10", "top 10"),
    ],
)
def test_label_normalisation(raw, expected):
    """A bare trailing number is part of the label, never a footnote marker."""
    assert normalise_label(raw) == expected


def test_cells_are_labelled_by_their_metric_and_period(clean_deck):
    """The rules read the figure index now, so this asserts what the index gives
    them: a metric from the column header and a period from the row label, on a
    table that puts the years down the side."""
    figures = [f for f in build_index(clean_deck) if f.source == "table"]
    assert figures, "the fixture has numeric tables"
    keyed = {(f.period, f.metric) for f in figures}
    assert ("FY2025A", "revenue") in keyed
    assert ("FY2025A", "ebitda") in keyed
    for figure in figures:
        assert figure.slide_index >= 1
        assert figure.reading is not None


def test_the_consistency_category_is_silent_on_the_clean_deck(
    clean_deck, reference_profile
):
    result = run_rules(clean_deck, reference_profile, include=["CO-*"])
    assert result.findings == [], "; ".join(
        f"slide {f.slide_index} {f.rule_id}: {f.message}" for f in result.findings
    )


def test_a_deck_with_no_tables_is_handled(tmp_path, reference_profile):
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = tmp_path / "empty.pptx"
    presentation.save(str(path))
    result = run_rules(load_deck(path), reference_profile, include=["CO-*"])
    assert result.findings == []


# --------------------------------------------------------------------------------------
# A figure stated more than twice, and figures stated on one slide
#
# Every test below reproduces a way the consistency rules used to fall silent on
# a real contradiction, or speak up about a correct table. They are the reason
# the module now classifies scale pair by pair, keys on the table rather than
# the slide, and derives its rounding tolerance per figure.
# --------------------------------------------------------------------------------------


def _deck_with_slides(path, slides):
    """Build a deck where each slide may carry several tables.

    ``slides`` is a list of slides, each a list of tables, each a list of rows.
    The single-table helper above cannot express two tables on one slide, which
    is precisely the case that went unchecked.
    """
    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    counter = 0
    for tables in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for position, rows in enumerate(tables):
            counter += 1
            frame = slide.shapes.add_table(
                len(rows),
                len(rows[0]),
                Pt(36),
                Pt(60 + position * 220),
                Pt(400),
                Pt(20 * len(rows)),
            )
            frame.name = f"Table {counter}"
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    frame.table.cell(row_index, column_index).text = str(value)
    presentation.save(str(path))
    return load_deck(path)


def _ebitda(value):
    return [["Fiscal year", "EBITDA"], ["2025A", value]]


def test_every_way_a_figure_is_stated_is_reported(tmp_path, reference_profile):
    """A figure stated three ways is two contradictions, not one.

    Only the first disagreement against the baseline was reported, so a reader
    told that 275 contradicts 263 was never told that 290 does too -- and would
    reconcile the pair they were shown and ship the deck.
    """
    deck = _deck_with_slides(
        tmp_path / "three.pptx",
        [[_ebitda("263")], [_ebitda("275")], [_ebitda("290")]],
    )
    result = _run(deck, reference_profile, "CO-001")

    assert len(result.findings) == 2, [f.message for f in result.findings]
    assert {f.slide_index for f in result.findings} == {2, 3}
    assert any("290" in f.message for f in result.findings)


def test_two_tables_on_one_slide_are_compared(tmp_path, reference_profile):
    """The guard excluded repetition inside one table by keying on the slide,
    which also excluded two different tables that happen to share one."""
    deck = _deck_with_slides(
        tmp_path / "same_slide.pptx", [[_ebitda("263"), _ebitda("275")]]
    )
    result = _run(deck, reference_profile, "CO-001")

    assert len(result.findings) == 1, [f.message for f in result.findings]
    assert "275" in result.findings[0].message


def test_one_units_error_does_not_excuse_a_real_contradiction(
    tmp_path, reference_profile
):
    """The scale test ran against the group's own minimum and maximum, so a
    figure restated in thousands anywhere under a label made every genuine
    disagreement under that label disappear. Each is now classified against the
    baseline on its own."""
    deck = _deck_with_slides(
        tmp_path / "masked.pptx",
        [[_ebitda("263")], [_ebitda("275")], [_ebitda("263000")]],
    )

    contradictions = _run(deck, reference_profile, "CO-001").findings
    assert len(contradictions) == 1, [f.message for f in contradictions]
    assert "275" in contradictions[0].message

    scales = _run(deck, reference_profile, "CO-002").findings
    assert len(scales) == 1, [f.message for f in scales]
    assert "263000" in scales[0].message


def test_a_correctly_rounded_column_of_mixed_precision_is_not_reported(
    tmp_path, reference_profile
):
    """10.04 + 20.4 + 30.4 is 60.84, and 60.8 is its correct one-decimal total.

    The tolerance took the column's *largest* precision and applied it to every
    addend, which is the tightest of the per-figure bounds rather than their
    sum, and reported correctly-rounded banking tables as not adding up.
    """
    deck = _deck_with_slides(
        tmp_path / "rounded.pptx",
        [
            [
                [
                    ["Segment", "FY25"],
                    ["North", "10.04"],
                    ["South", "20.4"],
                    ["East", "30.4"],
                    ["Total", "60.8"],
                ]
            ]
        ],
    )
    assert _run(deck, reference_profile, "CO-003").findings == []


def test_a_total_is_still_checked_when_a_figure_carries_a_footnote(
    tmp_path, reference_profile
):
    """One annotated cell used to disable the arithmetic for its whole column.

    10 + 20 + 30 + 5 is 65, not 99, and a marker on one of the addends is no
    reason to stop checking.
    """
    deck = _deck_with_slides(
        tmp_path / "footnote.pptx",
        [
            [
                [
                    ["Segment", "FY25"],
                    ["North", "10"],
                    ["South", "20"],
                    ["East", "30"],
                    ["West", "5 (a)"],
                    ["Total", "99"],
                ]
            ]
        ],
    )
    result = _run(deck, reference_profile, "CO-003")

    assert len(result.findings) == 1, [f.message for f in result.findings]
    assert "65" in result.findings[0].message


def test_a_bracketed_negative_is_not_mistaken_for_a_footnote(tmp_path):
    """The footnote stripper must not eat a parenthesised negative, which is how
    every banking table writes one."""
    from tieout.rules.consistency import read_cell_value, strip_value_footnote

    def value_of(text: str) -> float:
        reading = read_cell_value(text)
        assert reading is not None, f"{text!r} should read as a figure"
        return reading.value

    assert strip_value_footnote("(5)") == "(5)"
    assert value_of("(1,234)") == -1234.0
    assert value_of("(12)") == -12.0
    assert value_of("1,234 (2)") == 1234.0
    assert value_of("263*") == 263.0


def test_a_column_that_cannot_be_read_says_so(tmp_path, reference_profile):
    """A column the rule cannot add up and a column that adds up correctly were
    both silent. For a pre-send check those must never look the same."""
    deck = _deck_with_slides(
        tmp_path / "prose.pptx",
        [
            [
                [
                    ["Segment", "FY25"],
                    ["North", "10"],
                    ["South", "see note"],
                    ["East", "30"],
                    ["Total", "99"],
                ]
            ]
        ],
    )
    result = _run(deck, reference_profile, "CO-003")

    assert result.findings == []
    assert result.unchecked, "the column could not be read and nothing said so"
    assert "see note" in result.unchecked[0].reason


# --------------------------------------------------------------------------------------
# Past the tables
#
# The point of moving these rules onto tieout.figures, and PLAN.md section 3's
# first two rows. Before the index, CO-001 needed the figure in *two tables*: a
# headline was neither, and a chart carried no numeric values at all. Both of
# those are now the same rule doing the same thing to a wider set of figures,
# rather than a second implementation of what "the same figure" means.
# --------------------------------------------------------------------------------------


def _headline(path, table_rows, sentence):
    """One slide of table, one slide of prose that may or may not tie to it."""
    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = first.shapes.add_table(
        len(table_rows), len(table_rows[0]), Pt(36), Pt(120), Pt(500), Pt(20 * len(table_rows))
    )
    frame.name = "Table 1"
    for row_index, row in enumerate(table_rows):
        for column_index, value in enumerate(row):
            frame.table.cell(row_index, column_index).text = str(value)
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = second.shapes.add_textbox(Pt(36), Pt(60), Pt(800), Pt(40))
    box.text_frame.text = sentence
    presentation.save(str(path))
    return load_deck(path)


def test_co001_catches_a_headline_contradicting_its_table(tmp_path, reference_profile):
    """PLAN.md section 3, first row: "Revenue grew to $412m" over a table
    reading 408. CO-001 needed the figure in two tables; a headline is
    neither."""
    deck = _headline(
        tmp_path / "headline.pptx",
        _PROJECTIONS,
        "Revenue reached 2,100 in 2025A on continued depot expansion",
    )
    findings = findings_for(_run(deck, reference_profile, "CO-001"), "CO-001")
    assert findings, "the headline contradicts the table and must be reported"
    (finding,) = findings
    assert finding.slide_index == 2
    assert "2,100" in finding.message
    assert "1,908" in finding.message


def test_a_prose_finding_says_what_it_was_matched_on(tmp_path, reference_profile):
    """A prose anchor is much weaker evidence than a table label, and a reader
    deciding whether to act on the finding has to be able to see that."""
    deck = _headline(
        tmp_path / "evidence.pptx",
        _PROJECTIONS,
        "Revenue reached 2,100 in 2025A on continued depot expansion",
    )
    (finding,) = findings_for(_run(deck, reference_profile, "CO-001"), "CO-001")
    assert "matched on the sentence" in finding.message


def test_co001_is_silent_when_the_headline_agrees(tmp_path, reference_profile):
    """The half that matters. A rule that reported every sentence containing a
    number would be switched off by lunchtime."""
    deck = _headline(
        tmp_path / "agrees.pptx",
        _PROJECTIONS,
        "Revenue reached 1,908 in 2025A on continued depot expansion",
    )
    assert findings_for(_run(deck, reference_profile, "CO-001"), "CO-001") == []


def test_co001_is_silent_on_a_sentence_naming_no_metric_the_deck_uses(
    tmp_path, reference_profile
):
    deck = _headline(
        tmp_path / "ungrounded.pptx",
        _PROJECTIONS,
        "The team has grown to 240 people across the estate",
    )
    assert findings_for(_run(deck, reference_profile, "CO-001"), "CO-001") == []


def _table_and_chart(path, table_rows, categories, series):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = first.shapes.add_table(
        len(table_rows), len(table_rows[0]), Pt(36), Pt(120), Pt(500), Pt(20 * len(table_rows))
    )
    frame.name = "Table 1"
    for row_index, row in enumerate(table_rows):
        for column_index, value in enumerate(row):
            frame.table.cell(row_index, column_index).text = str(value)

    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = list(categories)
    for name, values in series:
        data.add_series(name, values)
    second.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(36), Pt(120), Pt(500), Pt(280), data
    )
    presentation.save(str(path))
    return load_deck(path)


def test_co001_catches_a_chart_contradicting_the_table(tmp_path, reference_profile):
    """PLAN.md section 3, second row: ChartModel carried no numeric values at
    all, so a chart contradicting the table beside it could not be seen."""
    deck = _table_and_chart(
        tmp_path / "chart.pptx",
        [["Fiscal year", "Revenue"], ["2024A", "1,562"], ["2025A", "1,908"]],
        ["2024A", "2025A"],
        [("Revenue", (1562.0, 2100.0))],
    )
    findings = findings_for(_run(deck, reference_profile, "CO-001"), "CO-001")
    assert findings
    (finding,) = findings
    assert finding.slide_index == 2
    assert "2,100" in finding.message


def test_co001_is_silent_when_the_chart_agrees(tmp_path, reference_profile):
    deck = _table_and_chart(
        tmp_path / "chart_agrees.pptx",
        [["Fiscal year", "Revenue"], ["2024A", "1,562"], ["2025A", "1,908"]],
        ["2024A", "2025A"],
        [("Revenue", (1562.0, 1908.0))],
    )
    assert findings_for(_run(deck, reference_profile, "CO-001"), "CO-001") == []


def test_a_cumulative_total_is_not_compared_against_a_single_year(
    tmp_path, reference_profile
):
    """The reference deck's own shape, and the false positive that closing the
    period range fixed. A "Total 2023A-2027E" row states five years of revenue;
    read as 2023A it contradicts the chart's first bar, and the clean deck
    reported itself."""
    deck = _table_and_chart(
        tmp_path / "cumulative.pptx",
        [
            ["Fiscal year", "Revenue"],
            ["2023A", "1,284"],
            ["2024A", "1,562"],
            ["Total 2023A-2024A", "2,846"],
        ],
        ["2023A", "2024A"],
        [("Revenue", (1284.0, 1562.0))],
    )
    assert findings_for(_run(deck, reference_profile, "CO-001"), "CO-001") == []


def test_a_figure_stated_in_two_scales_that_agree_is_not_a_contradiction(
    tmp_path, reference_profile
):
    """One table in millions and one in billions, both saying the same thing.
    Comparing the digits alone reports a factor of a thousand; comparing the
    figures reports nothing, which is correct."""
    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    for value, caption in (("1,908", "Figures in US$ millions."), ("1.908", "US$ bn")):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_table(2, 2, Pt(36), Pt(120), Pt(400), Pt(40))
        for row_index, row in enumerate([["Fiscal year", "Revenue"], ["2025A", value]]):
            for column_index, cell in enumerate(row):
                frame.table.cell(row_index, column_index).text = cell
        box = slide.shapes.add_textbox(Pt(36), Pt(420), Pt(400), Pt(20))
        box.text_frame.text = caption
    path = tmp_path / "scales.pptx"
    presentation.save(str(path))
    deck = load_deck(path)

    result = run_rules(deck, reference_profile, include=["CO-*"])
    assert result.findings == [], [f.message for f in result.findings]
