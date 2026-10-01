"""A corpus of decks with known defects and known intentional oddities.

PLAN.md §0. The suite proved for a long time that the tie-out is silent on a
clean deck, and nothing measured how often it is *wrong* on a deck with real
defects in it. A clean deck a rule is silent on proves nothing on its own (§8),
and a seeded defect a rule catches proves nothing about the deck where the same
rule cries wolf. This module holds both halves against the same rules and turns
them into two numbers per rule:

* **recall** -- of the defects a person would want reported, how many were;
* **false positives** -- findings on something labelled intentional, or on
  nothing labelled at all. In a deck whose every defect is labelled, an
  unlabelled finding is the tool disagreeing with the deck.

Every case is built here, from code, because client decks are never committed.
Each construction is one a real deck contains, and each label says why it is
labelled the way it is. Three of the intentional ones are the demo's own false
positives, verbatim; one of the defects is the case PLAN.md §0 predicts the
``kind`` dimension will make go quiet, kept in the corpus precisely so the
trade is measured rather than assumed.

Labels are deliberately on separate slides. Findings cluster per slide and per
rule, so two labels on one slide would let one finding satisfy both and inflate
recall.

``python -m tests.corpus`` prints the table PLAN.md §0.1 quotes.
"""

from __future__ import annotations

import sys
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from tests.test_real_world_constructions import _build as build_kestrel
from tests.test_real_world_constructions import _set_alpha
from tests.test_real_world_figures import (
    GOLD,
    GREY,
    NAVY,
    TIE_OUT_RULES,
    _blank,
    _chrome,
    _deck,
    _head,
    _rect,
    _table,
    _text,
)
from tests.test_real_world_figures import build as build_marlin
from tieout.fixtures.generator import build_all
from tieout.fixtures.spec import default_spec
from tieout.learn import learn_from_decks
from tieout.model.deck import DeckModel
from tieout.model.loader import load_deck
from tieout.profile.schema import Profile
from tieout.rules.base import AuditResult, Finding, clear_caches, run_rules

#: ``echo`` is a finding that is true but is another rule's seeded defect seen
#: from a second angle -- a nudged column that also crosses the margin. It is
#: neither a catch for this rule nor a false positive, and is not scored.
Truth = Literal["defect", "intentional", "echo"]

#: The layout rules whose false positives PLAN.md §0 is about. LO-002 is in
#: because it shares LO-001's bleed handling and would otherwise be where a
#: silenced overhang quietly re-appears.
LAYOUT_RULES: tuple[str, ...] = ("LO-001", "LO-002", "LO-003")


@dataclass(frozen=True)
class Label:
    """One thing in a deck that a rule must, or must not, report.

    ``slide`` is where the finding lands -- for a contradiction, the slide of
    the *later* statement, which is the one CO-001 reports. ``None`` matches
    anywhere, for a single-defect deck where the slide is not the point.
    ``shape`` narrows to one shape where a slide carries more than one
    candidate.
    """

    rule_id: str
    slide: int | None
    truth: Truth
    why: str
    shape: str | None = None

    def matches(self, finding: Finding) -> bool:
        if finding.rule_id != self.rule_id:
            return False
        if self.slide is not None and finding.slide_index != self.slide:
            return False
        return self.shape is None or finding.shape_name == self.shape


@dataclass(frozen=True)
class Case:
    name: str
    #: Builds the deck to check and the profile to check it against.
    build: Callable[[Path], tuple[DeckModel, Profile]]
    rules: tuple[str, ...]
    labels: tuple[Label, ...] = ()


@dataclass
class Score:
    """What one rule did across the corpus."""

    defects: int = 0
    caught: int = 0
    false_positives: int = 0
    missed: list[str] = field(default_factory=list)
    wrong: list[str] = field(default_factory=list)

    @property
    def recall(self) -> float:
        return self.caught / self.defects if self.defects else 1.0


@dataclass
class Report:
    scores: dict[str, Score]
    #: Everything a rule declined: refusals and, once intent evidence exists,
    #: positions it judged chosen. Kept beside the score because a refusal is
    #: not a pass (§8), and a recall figure that rose because a rule stopped
    #: looking is not a recall figure.
    declined: dict[str, list[str]]

    @property
    def defects(self) -> int:
        return sum(score.defects for score in self.scores.values())

    @property
    def caught(self) -> int:
        return sum(score.caught for score in self.scores.values())

    @property
    def false_positives(self) -> int:
        return sum(score.false_positives for score in self.scores.values())

    @property
    def recall(self) -> float:
        return self.caught / self.defects if self.defects else 1.0

    def table(self) -> str:
        lines = [f"{'rule':8} {'defects':>7} {'caught':>6} {'recall':>7} {'FP':>3}"]
        for rule_id in sorted(self.scores):
            score = self.scores[rule_id]
            lines.append(
                f"{rule_id:8} {score.defects:7} {score.caught:6} "
                f"{score.recall:7.0%} {score.false_positives:3}"
            )
        lines.append(
            f"{'all':8} {self.defects:7} {self.caught:6} {self.recall:7.0%} "
            f"{self.false_positives:3}"
        )
        for rule_id in sorted(self.scores):
            for line in self.scores[rule_id].missed:
                lines.append(f"  missed  {rule_id}: {line}")
            for line in self.scores[rule_id].wrong:
                lines.append(f"  wrong   {rule_id}: {line}")
        return "\n".join(lines)


def score(cases: Sequence[Case], directory: Path) -> Report:
    scores: dict[str, Score] = {}
    declined: dict[str, list[str]] = {}
    for case in cases:
        clear_caches()
        deck, profile = case.build(directory / case.name)
        result = run_rules(deck, profile, include=list(case.rules))
        if result.failed_rules:
            raise AssertionError(f"{case.name}: {result.failed_rules}")
        _score_case(case, result, scores)
        declined[case.name] = [
            f"{record.rule_id} slide {record.slide_index}: {record.reason}"
            for record in [*result.unchecked, *getattr(result, "excused", [])]
            if record.rule_id in case.rules
        ]
    return Report(scores=scores, declined=declined)


def _score_case(case: Case, result: AuditResult, scores: dict[str, Score]) -> None:
    for rule_id in case.rules:
        scores.setdefault(rule_id, Score())
    findings = [f for f in result.findings if f.rule_id in case.rules]
    claimed: set[int] = set()
    for label in case.labels:
        hits = [i for i, f in enumerate(findings) if label.matches(f)]
        score_ = scores.setdefault(label.rule_id, Score())
        if label.truth == "defect":
            score_.defects += 1
            if hits:
                score_.caught += 1
                claimed.update(hits)
            else:
                score_.missed.append(f"{case.name}: {label.why}")
    for position, finding in enumerate(findings):
        if position in claimed:
            continue
        if any(
            label.truth == "echo" and label.matches(finding) for label in case.labels
        ):
            continue
        intended = next(
            (
                label
                for label in case.labels
                if label.truth == "intentional" and label.matches(finding)
            ),
            None,
        )
        reason = f"(intentional: {intended.why})" if intended else "(unlabelled)"
        scores[finding.rule_id].false_positives += 1
        scores[finding.rule_id].wrong.append(
            f"{case.name} slide {finding.slide_index}: {finding.message[:110]} {reason}"
        )


# --------------------------------------------------------------------------------------
# The cases
# --------------------------------------------------------------------------------------


def _self_learned(deck_path: Path, client: str) -> tuple[DeckModel, Profile]:
    deck = load_deck(str(deck_path))
    return deck, learn_from_decks([deck], client).profile


def _marlin(defect: str | None) -> Callable[[Path], tuple[DeckModel, Profile]]:
    def build(directory: Path) -> tuple[DeckModel, Profile]:
        directory.mkdir(parents=True, exist_ok=True)
        return _self_learned(build_marlin(directory / "marlin.pptx", defect=defect), "marlin")

    return build


def _reference(which: Literal["clean", "dirty"]) -> Callable[[Path], tuple[DeckModel, Profile]]:
    """The generator's reference decks, checked against a profile learned from
    the clean one -- which is how a client is onboarded, rather than against the
    hand-built profile the unit tests use."""

    def build(directory: Path) -> tuple[DeckModel, Profile]:
        built = build_all(directory, default_spec())
        clean = load_deck(str(built.clean))
        profile = learn_from_decks([clean], "reference").profile
        clear_caches()
        return load_deck(str(getattr(built, which))), profile

    return build


def _reference_labels() -> tuple[Label, ...]:
    wanted = set(TIE_OUT_RULES) | set(LAYOUT_RULES)
    seeded = tuple(
        Label(d.rule_id, d.slide_index, "defect", f"seeded: {d.description}")
        for d in default_spec().defects
        if d.rule_id in wanted and d.variant is None
    )
    return seeded + REFERENCE_ECHOES


#: Every finding on the dirty reference deck that is not its rule's own seed,
#: traced to the seed that causes it. Each was checked by hand, once, against
#: :func:`tieout.fixtures.spec.default_defects`; a new unlabelled finding here
#: is scored as a false positive until someone does the same.
REFERENCE_ECHOES: tuple[Label, ...] = (
    Label("CO-001", 8, "echo",
          "TY-006 seeds 1,284.5 into slide 6's table; the chart on slide 8 plots 1,284, "
          "so the deck really does state FY2023A revenue two ways"),
    Label("LO-002", 11, "echo",
          "HY-008 replaces slide 11's logo with a low-resolution copy, which is then "
          "not recognised as the logo and is measured as content"),
    Label("LO-002", 15, "echo",
          "BR-006 replaces the page number with a non-conforming string, which is then "
          "not furniture and is measured as content"),
    Label("LO-002", 17, "echo", "BR-008 drags the title off its position and into the margin"),
    Label("LO-002", 24, "echo", "LO-003's nudged column crosses the margin by the same 3pt"),
)


# -- kinds: figures that share a metric and a quantity and are different facts ---------


def build_kinds(path: Path) -> Path:
    """Twelve slides of prose against one P&L, all of it correct except the
    four defects labelled in :data:`KINDS_LABELS`.

    Revenue 1,388 / 1,624 / 1,770 grows 17.0% then 9.0%; EBITDA 312 / 389 / 439
    gives margins of 22.5%, 24.0% and 24.8%, and grows 12.9% into FY25A.
    """
    presentation = _deck()
    cover = _blank(presentation)
    _text(cover, 36, 180, 600, 48, "Project Heron", size=38)
    _text(cover, 36, 240, 600, 24, "Management Presentation", size=14)
    _chrome(cover, 1)

    pnl = _blank(presentation)
    _head(pnl, "FINANCIAL PERFORMANCE", "Group Financial Performance")
    _table(pnl, 36, 120, 600, [
        ["$ in millions", "FY23A", "FY24A", "FY25A"],
        ["Revenue", "1,388", "1,624", "1,770"],
        ["EBITDA", "312", "389", "439"],
        ["EBITDA margin", "22.5%", "24.0%", "24.8%"],
        ["Churn", "4.6%", "4.1%", "3.8%"],
    ], "Group P&L")
    _chrome(pnl, 2)

    statements = (
        "64% of revenue is recurring in FY25A.",          # 3  share
        "9% revenue growth in FY25A.",                    # 4  change
        "EBITDA up 12.9% in FY25A.",                      # 5  change
        "EBITDA at 24.8% of revenue in FY25A.",           # 6  share of a base
        "Revenue up $146m in FY25A.",                     # 7  change, as an amount
        "Revenue grew 17% in FY24A.",                     # 8  change, correct
        "15% revenue growth in FY24A.",                   # 9  change, WRONG
        "EBITDA margin of 25.8% in FY25A.",               # 10 ratio, WRONG
        "Revenue of $1,790m in FY25A.",                   # 11 level, WRONG
        "Churn of 4.8% in FY25A.",                        # 12 no word says what kind
    )
    for page, sentence in enumerate(statements, start=3):
        slide = _blank(presentation)
        _head(slide, "HIGHLIGHTS", "Key Highlights")
        _text(slide, 36, 120, 880, 20, sentence)
        _chrome(slide, page)

    presentation.save(str(path))
    return path


KINDS_LABELS: tuple[Label, ...] = (
    Label("CO-001", 4, "intentional",
          "demo: '64% of revenue is recurring' and '9% revenue growth' are a share and a change"),
    Label("CO-001", 6, "intentional",
          "'EBITDA up 12.9%' and 'EBITDA at 24.8% of revenue' are a change and a share"),
    Label("CO-001", 7, "intentional",
          "'Revenue up $146m' is a change in revenue, not revenue"),
    Label("CO-001", 9, "defect",
          "'15% revenue growth in FY24A' contradicts 'Revenue grew 17% in FY24A'"),
    Label("CO-001", 10, "defect",
          "'EBITDA margin of 25.8%' contradicts the table's 24.8%"),
    Label("CO-001", 11, "defect",
          "'Revenue of $1,790m' contradicts the table's 1,770"),
    Label("CO-001", 12, "defect",
          "'Churn of 4.8%' contradicts the table's 3.8% -- no word says what kind of "
          "figure churn is, so PLAN.md §0 predicts this one goes quiet"),
)


def _kinds(directory: Path) -> tuple[DeckModel, Profile]:
    directory.mkdir(parents=True, exist_ok=True)
    return _self_learned(build_kinds(directory / "heron.pptx"), "heron")


# -- layout: a house style, and a later deck in it ------------------------------------

_OVAL = (760, -100, 374, 374)


def _osprey_chrome(slide, page: int) -> None:
    _text(slide, 790, 27, 130, 22, "OSPREY PARTNERS", size=6)
    _text(slide, 36, 510, 504, 14, "Project Osprey  |  Strictly Private", size=8, colour=GREY)
    _text(slide, 900, 510, 24, 14, str(page), size=8, colour=GREY)


def _osprey_content(
    presentation, page: int, title: str, *, right: float = 490, note: str | None = ""
):
    """A content slide: two columns of four points and a note across the foot.

    The note is what teaches the profile that content runs down to 450pt; a
    reference deck whose content stopped at 250pt learns a bottom margin there,
    and every later deck with a longer slide crosses it.
    """
    slide = _blank(presentation)
    _rect(slide, 36, 30, 4, 24, GOLD)
    _text(slide, 48, 36, 600, 24, title.upper(), size=14)
    _text(slide, 36, 66, 880, 22, f"{title} remains on plan.")
    for row in range(4):
        _text(slide, 36, 120 + row * 36, 426, 22,
              f"Left point {row + 1} on {title.lower()}.")
        box = _text(slide, right if row == 0 else 490, 120 + row * 36, 426, 22,
                    f"Right point {row + 1} on {title.lower()}.")
        if row == 0 and right != 490:
            box.name = "Body nudged"
    if note is not None:
        _text(slide, 36, 420, 880, 30, note or f"Note: {title.lower()} detail overleaf.")
    _text(slide, 36, 470, 400, 22, "Source: company information.", size=9, colour=GREY)
    _osprey_chrome(slide, page)
    return slide


def _bled_oval(slide) -> None:
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, *(Pt(v) for v in _OVAL))
    circle.fill.solid()
    circle.fill.fore_color.rgb = NAVY
    circle.line.fill.background()
    _set_alpha(circle, 12)
    # Sent to the back, where a cover device sits.
    tree = slide.shapes._spTree
    tree.remove(circle._element)
    tree.insert(2, circle._element)


def _osprey_divider(presentation, page: int, title: str, *, bleed: bool):
    slide = _blank(presentation)
    if bleed:
        _bled_oval(slide)
    _text(slide, 36, 220, 600, 48, title, size=32)
    _osprey_chrome(slide, page)
    return slide


def build_osprey_reference(path: Path) -> Path:
    """The deck the house style is learned from. Nothing in it bleeds."""
    presentation = _deck()
    cover = _blank(presentation)
    _text(cover, 36, 180, 600, 48, "Project Osprey", size=38)
    _text(cover, 36, 240, 600, 24, "Investor Presentation", size=14)
    _osprey_chrome(cover, 1)
    for page, title in enumerate(
        ("Market", "Company", "Strategy", "Customers", "Operations", "Financials",
         "Management", "Outlook"),
        start=2,
    ):
        _osprey_content(presentation, page, title)
    presentation.save(str(path))
    return path


def build_osprey_target(path: Path) -> Path:
    """A later deck in the same house style, with the labels in
    :data:`OSPREY_LABELS` built into it."""
    presentation = _deck()
    cover = _blank(presentation)
    _bled_oval(cover)                                              # 1
    _text(cover, 36, 180, 600, 48, "Project Osprey", size=38)
    _text(cover, 36, 240, 600, 24, "Investor Presentation", size=14)
    _osprey_chrome(cover, 1)

    _osprey_divider(presentation, 2, "Business overview", bleed=True)       # 2
    _osprey_content(presentation, 3, "Market", right=493)                    # 3

    dragged = _osprey_content(presentation, 4, "Company")                  # 4
    label = _text(dragged, 880, 300, 140, 22, "Indicative timetable, subject to change")
    label.name = "Dragged label"

    for page, title in ((5, "Strategy"), (6, "Customers"), (7, "Operations")):
        slide = _osprey_content(presentation, page, title, note=None)       # 5-7
        box = _text(slide, 493.5, 420, 410, 30,
                    f"Takeaway: {title.lower()} is a strength.", colour=GOLD)
        box.name = "Takeaway"

    _osprey_divider(presentation, 8, "Financial overview", bleed=True)     # 8

    panel_slide = _osprey_content(presentation, 9, "Financials")           # 9
    panel = panel_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(490), Pt(300), Pt(426), Pt(300))
    panel.fill.solid()
    panel.fill.fore_color.rgb = GREY
    panel.line.fill.background()
    _set_alpha(panel, 8)
    panel.name = "Stretched panel"
    tree = panel_slide.shapes._spTree
    tree.remove(panel._element)
    tree.insert(2, panel._element)

    presentation.save(str(path))
    return path


OSPREY_LABELS: tuple[Label, ...] = (
    Label("LO-001", 1, "intentional",
          "demo: a decorative oval bled off the cover on purpose"),
    Label("LO-001", 2, "intentional",
          "the same decorative oval, at the same place, on every divider"),
    Label("LO-001", 8, "intentional",
          "the same decorative oval, at the same place, on every divider"),
    Label("LO-003", 3, "defect",
          "a body text box nudged 3pt off the learned 490pt column", shape="Body nudged"),
    Label("LO-001", 4, "defect",
          "a text label dragged off the right edge", shape="Dragged label"),
    Label("LO-003", 5, "intentional",
          "demo: a takeaway box set 3.5pt inside the column on purpose, on three slides",
          shape="Takeaway"),
    Label("LO-003", 6, "intentional", "the same takeaway bar", shape="Takeaway"),
    Label("LO-003", 7, "intentional", "the same takeaway bar", shape="Takeaway"),
    Label("LO-001", 9, "defect",
          "a background panel stretched off the bottom of a content slide -- "
          "structurally identical to a decorative bleed", shape="Stretched panel"),
)


def _osprey(directory: Path) -> tuple[DeckModel, Profile]:
    directory.mkdir(parents=True, exist_ok=True)
    reference = load_deck(str(build_osprey_reference(directory / "reference.pptx")))
    profile = learn_from_decks([reference], "osprey").profile
    clear_caches()
    return load_deck(str(build_osprey_target(directory / "target.pptx"))), profile


def _kestrel(directory: Path) -> tuple[DeckModel, Profile]:
    directory.mkdir(parents=True, exist_ok=True)
    return _self_learned(build_kestrel(directory / "kestrel.pptx"), "kestrel")


def _seeded_marlin_label(rule_id: str) -> Label:
    return Label(rule_id, None, "defect", f"marlin seeded with one {rule_id} defect")


CASES: tuple[Case, ...] = (
    Case("marlin-clean", _marlin(None), TIE_OUT_RULES),
    *(
        Case(f"marlin-{rule_id}", _marlin(rule_id), (rule_id,), (_seeded_marlin_label(rule_id),))
        for rule_id in TIE_OUT_RULES
    ),
    Case("reference-clean", _reference("clean"), (*TIE_OUT_RULES, *LAYOUT_RULES)),
    Case("reference-dirty", _reference("dirty"), (*TIE_OUT_RULES, *LAYOUT_RULES),
         _reference_labels()),
    Case("heron-kinds", _kinds, TIE_OUT_RULES, KINDS_LABELS),
    Case("osprey-layout", _osprey, LAYOUT_RULES, OSPREY_LABELS),
    Case("kestrel-constructions", _kestrel, LAYOUT_RULES),
)


if __name__ == "__main__":  # pragma: no cover - the tool PLAN.md §0.1 is quoted from
    import tempfile

    with tempfile.TemporaryDirectory() as scratch:
        report = score(CASES, Path(scratch))
    print(report.table())
    if "-v" in sys.argv:
        for case, lines in report.declined.items():
            for line in lines:
                print(f"  declined {case}: {line}")
