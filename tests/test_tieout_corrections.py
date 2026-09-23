"""Acting on a tie-out finding: the Fix/Edit split, and the write path under it.

PLAN.md §5.5 draws one line and everything here holds it:

* **Fix it** where the figure is *derived* — a margin that must equal its
  inputs, a multiple, a column total, a bridge's close. The answer is
  arithmetic on figures the deck already prints, so TieOut is not choosing
  anything and the button writes.
* **Edit it** where two figures are each *stated* and disagree. Which is right
  is a judgement about the deal. TieOut opens the run with the counterpart
  beside it and writes nothing on its own.

Three things are tested, in that order of importance.

1. **The replacement is written in the original's own format.** A fix that
   corrects 24.0% to 19.0 and drops the per-cent sign has traded an arithmetic
   finding for a typography one, on a slide TieOut itself just edited. This is
   the detail §5.5 calls "most likely to be skipped and most likely to be
   noticed".
2. **The write lands in the right cell and nowhere else.** Positional, never a
   text match: two cells in one table can legitimately hold the same digits.
3. **A refusal is visible.** A chart point cannot be written, and the finding
   says so rather than offering a button that does nothing.
"""

from __future__ import annotations

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Pt

from tieout.figures import restate
from tieout.model.loader import load_deck
from tieout.profile.loader import PROFILE_DIR_ENV
from tieout.rules.base import clear_caches, run_rules
from tieout.text import parse_number
from tieout_fix import RECELL_KIND, apply_fix, plan_fixes, recell_fix
from tieout_ui.server import create_app
from tieout_ui.session import SessionStore

_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@pytest.fixture
def store():
    state = SessionStore()
    yield state
    state.close()


@pytest.fixture
def client(store, tmp_path, monkeypatch):
    monkeypatch.setenv(PROFILE_DIR_ENV, str(tmp_path / "profiles"))
    with TestClient(create_app(store)) as test_client:
        test_client.headers.update({"X-TieOut-Token": store.token})
        yield test_client


@pytest.fixture
def uploaded(client, clean_path):
    response = client.post(
        "/api/decks", files={"file": (clean_path.name, clean_path.read_bytes(), _MIME)}
    )
    assert response.status_code == 200, response.text
    return response.json()


@pytest.fixture
def onboarded(client, uploaded):
    response = client.post(
        "/api/learn", json={"deck_id": uploaded["deck_id"], "client": "demo"}
    )
    assert response.status_code == 200, response.text
    return response.json()


def _deck(path, slides):
    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    for index, rows in enumerate(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_table(
            len(rows), len(rows[0]), Pt(36), Pt(120), Pt(600), Pt(20 * len(rows))
        )
        frame.name = f"Table {index + 1}"
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                frame.table.cell(row_index, column_index).text = str(value)
    presentation.save(str(path))
    return load_deck(path)


def _cells(path):
    """Every cell of every table in a saved deck, as text."""
    out = []
    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            if shape.has_table:
                out.append(
                    [[cell.text for cell in row.cells] for row in shape.table.rows]
                )
    return out


def _findings(deck, profile, rule_id):
    return [
        f
        for f in run_rules(deck, profile, include=[rule_id]).findings
        if f.rule_id == rule_id
    ]


# --------------------------------------------------------------------------------------
# The replacement keeps the original's format
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("original", "value", "expected"),
    [
        ("1,908", 2100.0, "2,100"),
        ("1908", 2100.0, "2100"),
        ("18.4%", 19.0164, "19.0%"),
        ("9.2x", 8.1977, "8.2x"),
        ("8,420", 9000.0, "9,000"),
        ("351", -12.0, "-12"),
        ("US$8,420", 9000.4, "US$9,000"),
        ("£12.50", 13.257, "£13.26"),
        ("0.0%", 100.0, "100.0%"),
    ],
)
def test_a_replacement_is_written_in_the_originals_format(original, value, expected):
    reading = parse_number(original)
    assert reading is not None, original
    assert restate(reading, value) == expected


def test_a_parenthesised_negative_stays_parenthesised(tmp_path):
    reading = parse_number("(42)")
    assert reading is not None
    assert restate(reading, -55.0) == "(55)"


def test_a_parenthesised_figure_turning_positive_drops_the_brackets():
    """Writing "(17)" for a positive number would be worse than either choice:
    the brackets mean negative, and keeping them would make the fix lie."""
    reading = parse_number("(42)")
    assert reading is not None
    assert restate(reading, 17.0) == "17"


def test_the_separator_the_deck_uses_is_the_one_written_back():
    reading = parse_number("1 234")
    assert reading is not None
    assert restate(reading, 5678.0) == "5 678"


def test_a_corrected_figure_does_not_introduce_a_typography_finding(
    tmp_path, reference_profile
):
    """The end-to-end version of the point, and the one that matters.

    TY-006 reports a figure whose decimals do not match its column. A fix that
    corrected the arithmetic and wrote a bare "19" into a column of one-decimal
    percentages would clear CO-004 and raise TY-006 in the same breath.
    """
    deck = _deck(
        tmp_path / "format.pptx",
        [[
            ["Fiscal year", "Revenue", "EBITDA", "Margin"],
            ["2024A", "1,562", "263", "16.8%"],
            ["2025A", "1,908", "351", "24.0%"],
        ]],
    )
    clear_caches()
    fixes = plan_fixes(run_rules(deck, reference_profile, include=["CO-004"]), reference_profile)
    (fix,) = fixes.values()
    out = tmp_path / "fixed.pptx"
    assert apply_fix(tmp_path / "format.pptx", out, fix).applied

    (table,) = _cells(out)
    assert table[2][3] == "18.4%", "the per-cent sign and the decimal must survive"

    clear_caches()
    after = run_rules(load_deck(out), reference_profile, include=["CO-*", "TY-006"])
    assert after.findings == [], [f.message for f in after.findings]


# --------------------------------------------------------------------------------------
# The write lands where it is addressed
# --------------------------------------------------------------------------------------


def test_a_cell_fix_changes_one_cell_and_leaves_the_rest(tmp_path, reference_profile):
    rows = [
        ["Company", "Enterprise value", "EBITDA", "Multiple"],
        ["Calderwood Logistics", "8,420", "912", "9.2x"],
        ["Pemberton Freight", "6,180", "674", "11.5x"],
        ["Thorne Distribution", "4,935", "602", "8.2x"],
    ]
    source = tmp_path / "comps.pptx"
    deck = _deck(source, [rows])
    clear_caches()
    fixes = plan_fixes(run_rules(deck, reference_profile, include=["CO-005"]), reference_profile)
    (fix,) = fixes.values()
    out = tmp_path / "comps-fixed.pptx"
    assert apply_fix(source, out, fix).applied

    (table,) = _cells(out)
    assert table[2][3] == "9.2x", "Pemberton's multiple is 6,180 over 674"
    for row in (0, 1, 3):
        assert table[row] == rows[row], f"row {row} must be untouched"


def test_the_same_digits_in_another_cell_are_not_rewritten(tmp_path, reference_profile):
    """Positional, never a text match. Two rows here both read 9.2x and only
    one of them is wrong; a substring rewrite would correct both."""
    rows = [
        ["Company", "Enterprise value", "EBITDA", "Multiple"],
        ["Calderwood Logistics", "8,420", "912", "9.2x"],
        ["Ellesmere Transport", "2,486", "271", "9.2x"],
        ["Ravensworth Group", "3,710", "387", "9.2x"],
    ]
    source = tmp_path / "same.pptx"
    deck = _deck(source, [rows])
    clear_caches()
    fixes = plan_fixes(run_rules(deck, reference_profile, include=["CO-005"]), reference_profile)
    assert fixes, "Ravensworth is 9.6x, not 9.2x"
    (fix,) = fixes.values()
    out = tmp_path / "same-fixed.pptx"
    assert apply_fix(source, out, fix).applied

    (table,) = _cells(out)
    assert [row[3] for row in table] == ["Multiple", "9.2x", "9.2x", "9.6x"]


def test_a_cell_fix_is_addressed_by_row_and_column(tmp_path):
    """Asserted on the fix itself, not only on the outcome: a payload that
    happened to be right for a one-row table would hide an off-by-one."""
    fix = recell_fix(
        slide_index=1, shape_id=7, shape_name="Table 1",
        row=3, column=2, paragraph=0, run=0, text="9.6x",
    )
    assert fix.kind == RECELL_KIND
    assert fix.payload["row"] == 3
    assert fix.payload["column"] == 2
    assert "9.6x" in fix.summary


def test_a_cell_fix_refuses_an_address_the_table_does_not_have(tmp_path):
    """The realistic way to get here is a deck edited under a page still
    showing the previous check. It must refuse, and leave the file alone."""
    source = tmp_path / "small.pptx"
    _deck(source, [[["Fiscal year", "Revenue"], ["2025A", "1,908"]]])
    before = source.read_bytes()
    out = tmp_path / "out.pptx"
    report = apply_fix(
        source,
        out,
        recell_fix(
            slide_index=1, shape_id=_first_shape_id(source), shape_name="Table 1",
            row=9, column=9, paragraph=0, run=0, text="nope",
        ),
    )
    assert not report.applied
    assert "row 9" in report.detail
    assert out.read_bytes() == before, "a refused fix must leave the deck as it was"


def _first_shape_id(path):
    slide = Presentation(str(path)).slides[0]
    return slide.shapes[0].shape_id


def test_a_cell_fix_leaves_every_other_part_of_the_package_alone(
    tmp_path, reference_profile
):
    source = tmp_path / "pkg.pptx"
    deck = _deck(
        source,
        [[
            ["Fiscal year", "Revenue", "EBITDA", "Margin"],
            ["2025A", "1,908", "351", "24.0%"],
        ]],
    )
    clear_caches()
    (fix,) = plan_fixes(
        run_rules(deck, reference_profile, include=["CO-004"]), reference_profile
    ).values()
    out = tmp_path / "pkg-fixed.pptx"
    assert apply_fix(source, out, fix).applied

    import zipfile

    with zipfile.ZipFile(source) as a, zipfile.ZipFile(out) as b:
        assert set(a.namelist()) == set(b.namelist())
        changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
    assert len(changed) == 1, f"one part should change, not {changed}"
    assert changed[0].startswith("ppt/slides/slide")


# --------------------------------------------------------------------------------------
# Fix, Edit, and refusal
# --------------------------------------------------------------------------------------


def test_a_derived_finding_carries_a_fix(tmp_path, reference_profile):
    deck = _deck(
        tmp_path / "derived.pptx",
        [[
            ["Fiscal year", "Revenue", "EBITDA", "Margin"],
            ["2025A", "1,908", "351", "24.0%"],
        ]],
    )
    (finding,) = _findings(deck, reference_profile, "CO-004")
    assert finding.correction is not None
    assert finding.correction.kind == "fix"
    assert finding.correction.replacement == "18.4%"
    assert finding.correction.writable


def test_a_contradiction_between_two_stated_figures_carries_an_edit(
    tmp_path, reference_profile
):
    """No replacement, and the counterpart named. TieOut can see the two
    figures disagree and has no way to know which the deal supports."""
    deck = _deck(
        tmp_path / "stated.pptx",
        [
            [["Fiscal year", "EBITDA"], ["2025A", "351"]],
            [["Fiscal year", "EBITDA"], ["2025A", "375"]],
        ],
    )
    (finding,) = _findings(deck, reference_profile, "CO-001")
    correction = finding.correction
    assert correction is not None
    assert correction.kind == "edit"
    assert correction.replacement is None
    assert correction.counterpart == "351"
    assert correction.counterpart_slide == 1


def test_no_fix_is_planned_for_a_contradiction(tmp_path, reference_profile):
    """The other half, asserted where it counts: on what the buttons offer."""
    deck = _deck(
        tmp_path / "nofix.pptx",
        [
            [["Fiscal year", "EBITDA"], ["2025A", "351"]],
            [["Fiscal year", "EBITDA"], ["2025A", "375"]],
        ],
    )
    clear_caches()
    fixes = plan_fixes(
        run_rules(deck, reference_profile, include=["CO-001"]), reference_profile
    )
    assert fixes == {}


def test_a_chart_anchored_finding_refuses_visibly(tmp_path, reference_profile):
    """PLAN.md §6: the write path that does not exist "must refuse, visibly,
    rather than appearing to work". A chart's values live in a cached copy and
    in an embedded workbook, and TieOut writes neither."""
    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = first.shapes.add_table(3, 2, Pt(36), Pt(120), Pt(400), Pt(60))
    for r, row in enumerate([["Fiscal year", "Revenue"], ["2024A", "1,562"], ["2025A", "1,908"]]):
        for c, value in enumerate(row):
            frame.table.cell(r, c).text = value
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["2024A", "2025A"]
    data.add_series("Revenue", (1562.0, 2100.0))
    second.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(36), Pt(120), Pt(500), Pt(280), data
    )
    path = tmp_path / "chart.pptx"
    presentation.save(str(path))

    (finding,) = _findings(load_deck(path), reference_profile, "CO-001")
    correction = finding.correction
    assert correction is not None
    assert correction.source == "chart"
    assert not correction.writable
    assert "chart" in (correction.refused or "")


# --------------------------------------------------------------------------------------
# Over HTTP
# --------------------------------------------------------------------------------------


def test_the_audit_view_carries_the_correction(client, onboarded, uploaded):
    body = client.post(
        "/api/check", json={"deck_id": uploaded["deck_id"], "client": "demo"}
    ).json()
    for slide in body["slides"]:
        for finding in slide["findings"]:
            assert "figure" in finding, finding["rule_id"]


def test_editing_a_table_cell_over_http(client, onboarded, dirty_path, store):
    """The endpoint the **Edit it** button posts to. Before the cell address
    existed it answered "that shape has no text frame" -- on the one kind of
    shape every tie-out finding is about."""
    upload = client.post(
        "/api/decks", files={"file": (dirty_path.name, dirty_path.read_bytes(), _MIME)}
    ).json()
    body = client.post(
        "/api/check", json={"deck_id": upload["deck_id"], "client": "demo"}
    ).json()

    target = next(
        (finding, slide["index"])
        for slide in body["slides"]
        for finding in slide["findings"]
        if finding.get("figure")
        and finding["figure"]["source"] == "table"
        and not finding["figure"]["refused"]
    )
    finding, slide_index = target
    figure = finding["figure"]

    response = client.post(
        "/api/edit-text",
        json={
            "deck_id": upload["deck_id"],
            "client": "demo",
            "slide": slide_index,
            "shape_id": figure["shape_id"],
            "row": figure["row"],
            "column": figure["column"],
            "paragraph": figure["paragraph"],
            "run": figure["run"],
            "text": "1,234",
        },
    )
    assert response.status_code == 200, response.text
    assert "edited_text" in response.json()

    deck = store.get(upload["deck_id"])
    written = _cells(deck.current)
    assert any("1,234" in cell for table in written for row in table for cell in row)


def test_editing_a_cell_that_does_not_exist_is_refused_with_a_sentence(
    client, onboarded, dirty_path
):
    """The clean deck has no tie-out finding to borrow an address from -- which
    is the point of it -- so this runs against the seeded one."""
    upload = client.post(
        "/api/decks", files={"file": (dirty_path.name, dirty_path.read_bytes(), _MIME)}
    ).json()
    body = client.post(
        "/api/check", json={"deck_id": upload["deck_id"], "client": "demo"}
    ).json()
    table = next(
        (finding, slide["index"])
        for slide in body["slides"]
        for finding in slide["findings"]
        if finding.get("figure") and finding["figure"]["source"] == "table"
    )
    finding, slide_index = table
    response = client.post(
        "/api/edit-text",
        json={
            "deck_id": upload["deck_id"],
            "client": "demo",
            "slide": slide_index,
            "shape_id": finding["figure"]["shape_id"],
            "row": 99,
            "column": 99,
            "paragraph": 0,
            "run": 0,
            "text": "x",
        },
    )
    assert response.status_code == 422
    assert "row 100" in response.json()["detail"]


def test_the_page_offers_edit_where_the_figures_are_stated(client):
    """Asserted on the served page: a server that marks a finding editable and
    a page with no button for it is a decision nobody can act on."""
    page = client.get("/").text
    assert "data-editfig" in page
    assert "Edit it" in page
    assert "showCounterpart" in page
    assert "cell_text_editable" in page


def test_a_table_cell_is_drawn_as_addressable_runs(client, onboarded, uploaded):
    """The surface has to hold the address a correction names, or **Edit it**
    has nothing to open."""
    body = client.get(
        f"/api/canvas/{uploaded['deck_id']}/6", params={"client": "demo"}
    )
    if body.status_code == 404:  # pragma: no cover - depends on the fixture deck
        pytest.skip("slide 6 is not in this fixture")
    view = body.json()

    def walk(nodes):
        for node in nodes:
            yield node
            yield from walk(node.get("children", []))

    tables = [n for n in walk(view["shapes"]) if n.get("table")]
    assert tables, "slide 6 of the reference deck is a table"
    for node in tables:
        assert node["editable"]["cell_text_editable"]
        cell = next(c for c in node["table"]["cells"] if not c["is_merge_continuation"])
        assert cell["paragraphs"], "a cell must expose its paragraphs"


def test_the_canvas_still_refuses_a_chart(client, onboarded, uploaded):
    view = client.get(
        f"/api/canvas/{uploaded['deck_id']}/8", params={"client": "demo"}
    ).json()

    def walk(nodes):
        for node in nodes:
            yield node
            yield from walk(node.get("children", []))

    charts = [n for n in walk(view["shapes"]) if n["kind"] == "chart"]
    assert charts, "slide 8 of the reference deck is a chart"
    for node in charts:
        assert not node["editable"]["cell_text_editable"]
        assert not node["editable"]["text_editable"]


def test_a_derived_fix_applies_through_the_ordinary_fix_button(
    client, onboarded, dirty_path, store
):
    """No new endpoint. A derived correction is an ordinary fix with an
    ordinary key, so the button that applies a recolour applies this too."""
    upload = client.post(
        "/api/decks", files={"file": (dirty_path.name, dirty_path.read_bytes(), _MIME)}
    ).json()
    body = client.post(
        "/api/check", json={"deck_id": upload["deck_id"], "client": "demo"}
    ).json()
    action = next(
        a for a in body["actions"]
        if a["rule_id"] in ("CO-003", "CO-004", "CO-005", "CO-006", "CO-007")
        and a["fixable"]
    )
    before = len(store.get(upload["deck_id"]).history)
    response = client.post(
        "/api/fix",
        json={"deck_id": upload["deck_id"], "client": "demo", "key": action["key"]},
    )
    assert response.status_code == 200, response.text
    assert len(store.get(upload["deck_id"]).history) == before + 1


def test_a_fix_can_be_undone(client, onboarded, dirty_path, store):
    upload = client.post(
        "/api/decks", files={"file": (dirty_path.name, dirty_path.read_bytes(), _MIME)}
    ).json()
    body = client.post(
        "/api/check", json={"deck_id": upload["deck_id"], "client": "demo"}
    ).json()
    action = next(
        a for a in body["actions"]
        if a["rule_id"] in ("CO-003", "CO-004", "CO-005", "CO-006", "CO-007")
        and a["fixable"]
    )
    deck = store.get(upload["deck_id"])
    original = deck.current.read_bytes()

    client.post(
        "/api/fix",
        json={"deck_id": upload["deck_id"], "client": "demo", "key": action["key"]},
    )
    assert deck.current.read_bytes() != original

    client.post("/api/undo", json={"deck_id": upload["deck_id"], "client": "demo"})
    assert deck.current.read_bytes() == original
