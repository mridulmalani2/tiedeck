"""Layout rule tests.

Two tests per rule: it catches its own seeded defect, and it says nothing about
the clean deck. The second matters more. A layout rule that fires on correct
slides is worse than no rule at all, because a report full of false alignment
findings teaches the reader to skim past the real ones.

The dirty deck deliberately carries every co-existable defect at once, so a
layout rule legitimately reports collateral damage from another rule's seed --
the off-canvas box seeded for LO-001 also intrudes into the safe margin, and the
unapproved-typeface box seeded for BR-005 overlaps its neighbour. So detection
tests assert the seeded slide is *among* the findings rather than the only one,
and the clean-deck tests carry the burden of proving there is no noise.
"""

from __future__ import annotations

import pytest

from tests.conftest import assert_silent_on_clean, findings_for, slide_indices
from tieout.fixtures.spec import default_defects
from tieout.model.furniture import detect_furniture
from tieout.profile.schema import FontRole
from tieout.rules.base import clear_caches, run_rules

SEEDED = {d.rule_id: d.slide_index for d in default_defects() if d.variant is None}

LAYOUT_RULE_IDS = (
    "LO-001",
    "LO-002",
    "LO-003",
    "LO-004",
    "LO-005",
    "LO-006",
    "LO-007",
    "LO-008",
)


def _run(deck, profile, rule_id):
    return run_rules(deck, profile, include=[rule_id])


def _assert_catches(deck, profile, rule_id):
    result = _run(deck, profile, rule_id)
    assert rule_id in result.rules_run, (
        f"{rule_id} did not run: {[(s.rule_id, s.reason) for s in result.rules_skipped]}"
    )
    findings = findings_for(result, rule_id)
    assert SEEDED[rule_id] in slide_indices(findings), (
        f"{rule_id} missed its seeded defect on slide {SEEDED[rule_id]}; "
        f"reported {sorted(slide_indices(findings))}"
    )
    return findings


# --------------------------------------------------------------------------------------
# Registration
# --------------------------------------------------------------------------------------


def test_every_layout_rule_is_registered():
    from tieout.rules.base import load_all_rules

    registry = load_all_rules()
    for rule_id in LAYOUT_RULE_IDS:
        assert rule_id in registry, f"{rule_id} is not registered"
        assert registry[rule_id].category == "layout"


def test_every_layout_rule_documents_itself():
    """Section 9 requires each rule to state what it measures and how it can be
    wrong. ``tieout rules`` surfaces both, so an undocumented rule is a gap the
    user sees."""
    from tieout.rules.base import load_all_rules

    registry = load_all_rules()
    for rule_id in LAYOUT_RULE_IDS:
        rule = registry[rule_id]
        assert rule.__doc__ and len(rule.__doc__.strip()) > 80, rule_id
        assert rule.summary, rule_id


# --------------------------------------------------------------------------------------
# LO-001 off canvas
# --------------------------------------------------------------------------------------


def test_lo001_catches_a_shape_off_the_canvas(dirty_deck, reference_profile):
    findings = _assert_catches(dirty_deck, reference_profile, "LO-001")
    seeded = [f for f in findings if f.slide_index == SEEDED["LO-001"]]
    assert any(f.measured for f in seeded), "the overhang should be measured in points"


def test_lo001_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    assert_silent_on_clean(_run(clean_deck, reference_profile, "LO-001"), "LO-001")


def test_lo001_accounts_for_rotation(clean_deck, reference_profile):
    """A rotated shape can overhang the canvas while its stored geometry sits
    inside it, so the rule must measure the rotated bounding box."""
    slide = clean_deck.slides[3]
    shape = next(s for s in slide.leaf_shapes() if s.width_pt > 200)
    original_rotation, original_left = shape.rotation, shape.left_pt
    try:
        shape.left_pt = clean_deck.width_pt - shape.width_pt - 2
        shape.rotation = 45.0
        findings = findings_for(_run(clean_deck, reference_profile, "LO-001"), "LO-001")
        assert slide.index in slide_indices(findings)
    finally:
        shape.rotation, shape.left_pt = original_rotation, original_left


def _deck_with_a_bled_graphic(path):
    """One slide: a decorative graphic off the corner, content in front of it."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Pt

    from tieout.model.loader import load_deck

    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(760), Pt(-100), Pt(374), Pt(374))
    oval.fill.solid()
    oval.line.fill.background()
    for index in range(3):
        box = slide.shapes.add_textbox(Pt(36), Pt(100 + index * 40), Pt(400), Pt(30))
        box.name = f"Body {index + 1}"
        box.text_frame.text = f"Line {index + 1} of the slide's own content"
    presentation.save(str(path))
    return load_deck(path)


def test_lo001_reports_a_bleed_where_the_reference_deck_had_none(
    tmp_path, reference_profile
):
    """No evidence either way, so the info line stays: this is the old behaviour,
    and it is what a house style that does not bleed should still get."""
    deck = _deck_with_a_bled_graphic(tmp_path / "bleed-unseen.pptx")
    profile = reference_profile.model_copy(deep=True)
    profile.layout.decorative_bleed_slides = []
    clear_caches()
    findings = findings_for(_run(deck, profile, "LO-001"), "LO-001")
    assert [f.severity for f in findings] == ["info"]


def test_lo001_says_nothing_where_the_reference_deck_bled_too(
    tmp_path, reference_profile
):
    """The reference deck bled, so bleeding is house style and not a finding.

    Before this, a profile learned from a deck with a cover device reported that
    same device on that same deck, with the remedy "Nothing to do unless this was
    not intended".
    """
    deck = _deck_with_a_bled_graphic(tmp_path / "bleed-known.pptx")
    profile = reference_profile.model_copy(deep=True)
    profile.layout.decorative_bleed_slides = [1]
    clear_caches()
    assert not findings_for(_run(deck, profile, "LO-001"), "LO-001")


# --------------------------------------------------------------------------------------
# LO-002 margin intrusion
# --------------------------------------------------------------------------------------


def test_lo002_catches_a_margin_intrusion(dirty_deck, reference_profile):
    _assert_catches(dirty_deck, reference_profile, "LO-002")


def test_lo002_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    assert_silent_on_clean(_run(clean_deck, reference_profile, "LO-002"), "LO-002")


def test_lo002_is_skipped_without_learned_margins(dirty_deck, reference_profile):
    reference_profile.layout.safe_margin_pt = {}
    result = _run(dirty_deck, reference_profile, "LO-002")
    assert result.findings == []
    assert any(s.rule_id == "LO-002" for s in result.rules_skipped)


def test_lo002_ignores_archetypes_with_no_learned_margin(dirty_deck, reference_profile):
    """An archetype absent from the profile is unchecked, not assumed."""
    reference_profile.layout.safe_margin_pt.pop("content", None)
    findings = findings_for(_run(dirty_deck, reference_profile, "LO-002"), "LO-002")
    assert SEEDED["LO-002"] not in slide_indices(findings)


def test_one_misplaced_shape_is_not_reported_by_two_rules(dirty_deck, reference_profile):
    """A shape off the canvas is off the safe margin too, necessarily.

    The margin lies inside the canvas, so LO-002 restating LO-001 is one
    misplacement measured twice -- and measured differently, because LO-001
    bounds the frame and LO-002 the ink. On the defect deck the seeded off-canvas
    box came back as "extends off the canvas: right by 120.0pt" *and* "intrudes
    into the safe margin: right by ...", which reads as two things to fix.

    Nothing is lost by saying it once: the shape has to come back onto the slide
    either way, and the margin is measured again on the next run.
    """
    clear_caches()
    result = run_rules(dirty_deck, reference_profile, include=["LO-001", "LO-002"])
    seen: dict[tuple[int, object], set[str]] = {}
    for finding in result.findings:
        key = (finding.slide_index, getattr(finding.where, "name", None))
        seen.setdefault(key, set()).add(finding.rule_id)
    doubled = {key: rules for key, rules in seen.items() if len(rules) > 1}
    assert not doubled, f"one shape reported by two rules: {doubled}"


# --------------------------------------------------------------------------------------
# LO-003 near-miss alignment
# --------------------------------------------------------------------------------------


def test_lo003_catches_a_column_off_its_grid_line(dirty_deck, reference_profile):
    _assert_catches(dirty_deck, reference_profile, "LO-003")


def test_lo003_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    """The whole reason the fixture is built on a 12pt grid.

    Every shape edge either coincides exactly or differs by at least 12pt, so
    there is no near miss to find. If this fails, the rule is over-reporting.
    """
    assert_silent_on_clean(_run(clean_deck, reference_profile, "LO-003"), "LO-003")


def test_lo003_needs_a_grid_to_judge_against(dirty_deck, reference_profile):
    """Without a learned grid a 3pt offset is indistinguishable from a deliberate
    one, which is section 8.3's point: off a real grid line it is a defect, off a
    one-off edge it is not."""
    reference_profile.layout.grid.columns_pt = []
    reference_profile.layout.grid.rows_pt = []
    result = _run(dirty_deck, reference_profile, "LO-003")
    assert findings_for(result, "LO-003") == []


def test_lo003_clusters_rather_than_reporting_every_pair(dirty_deck, reference_profile):
    """A row of misaligned blocks is one problem, not one per pair."""
    findings = findings_for(_run(dirty_deck, reference_profile, "LO-003"), "LO-003")
    per_slide = dict.fromkeys(slide_indices(findings), 0)
    for finding in findings:
        per_slide[finding.slide_index] += 1
    assert all(count <= 4 for count in per_slide.values()), per_slide


# --------------------------------------------------------------------------------------
# LO-004 overlap
# --------------------------------------------------------------------------------------


def test_lo004_catches_overlapping_text_shapes(dirty_deck, reference_profile):
    _assert_catches(dirty_deck, reference_profile, "LO-004")


def test_lo004_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    assert_silent_on_clean(_run(clean_deck, reference_profile, "LO-004"), "LO-004")


def test_lo004_tolerates_a_rule_beneath_text(clean_deck, reference_profile):
    """Hairline rules legitimately underlie text and must not be reported.

    The clean deck puts a 1.5pt navy rule directly beneath every title, so this
    is already covered by the silence test; asserting it explicitly documents the
    exemption so a future change to it is deliberate.
    """
    thin = [
        shape
        for slide in clean_deck.slides
        for shape in slide.leaf_shapes()
        if 0 < shape.height_pt < 3
    ]
    assert thin, "the fixture is supposed to contain hairline rules"
    assert findings_for(_run(clean_deck, reference_profile, "LO-004"), "LO-004") == []


# --------------------------------------------------------------------------------------
# LO-005 displaced recurring element
# --------------------------------------------------------------------------------------


def test_lo005_catches_a_displaced_footnote(dirty_deck, reference_profile):
    _assert_catches(dirty_deck, reference_profile, "LO-005")


def test_lo005_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    assert_silent_on_clean(_run(clean_deck, reference_profile, "LO-005"), "LO-005")


def test_lo005_is_skipped_without_learned_recurring_elements(
    dirty_deck, reference_profile
):
    reference_profile.layout.recurring = []
    result = _run(dirty_deck, reference_profile, "LO-005")
    assert result.findings == []
    assert any(s.rule_id == "LO-005" for s in result.rules_skipped)


# --------------------------------------------------------------------------------------
# LO-006 text overflow
# --------------------------------------------------------------------------------------


def test_lo006_is_disabled_by_default(clean_deck, reference_profile):
    """A glob selection is a filter, not an opt-in, so LO-006 stays off.

    Naming the rule exactly is a request to run it, which is the only way to
    reach an off-by-default rule without editing the profile; that path is
    covered by the detection test below.
    """
    result = run_rules(clean_deck, reference_profile, include=["LO-0*"])
    assert "LO-006" not in result.rules_run
    assert any(
        s.rule_id == "LO-006" and "default" in s.reason for s in result.rules_skipped
    )


def test_lo006_can_be_opted_into_through_the_profile(clean_deck, reference_profile):
    reference_profile.rules.enabled = ["LO-006"]
    result = run_rules(clean_deck, reference_profile, include=["LO-0*"])
    assert "LO-006" in result.rules_run


def test_lo006_either_reports_or_records_the_seeded_overflow(
    dirty_deck, reference_profile
):
    """LO-006 needs the real font file and must never guess.

    The seeded defect is set in the figure typeface, which resolves on this
    platform through a metric-compatible substitute, so a finding is expected.
    But the rule is specified to skip a shape whose font cannot be resolved and
    record it in ``unchecked`` instead, and on a host with no matching font that
    is the correct outcome. Both are accepted; what is not accepted is silence.
    """
    reference_profile.rules.disabled = []
    result = run_rules(dirty_deck, reference_profile, include=["LO-006"])
    seeded = SEEDED["LO-006"]
    reported = seeded in slide_indices(findings_for(result, "LO-006"))
    recorded = any(
        u.rule_id == "LO-006" and u.slide_index == seeded for u in result.unchecked
    )
    assert reported or recorded, (
        f"LO-006 neither reported nor recorded the overflow seeded on slide {seeded}"
    )


def test_lo006_never_reports_overflow_on_the_clean_deck(clean_deck, reference_profile):
    """Unchecked entries are fine here; findings are not."""
    reference_profile.rules.disabled = []
    result = run_rules(clean_deck, reference_profile, include=["LO-006"])
    assert findings_for(result, "LO-006") == [], [
        f.message for f in findings_for(result, "LO-006")
    ]


def test_lo006_records_unresolvable_fonts_rather_than_guessing(
    clean_deck, reference_profile
):
    reference_profile.rules.disabled = []
    result = run_rules(clean_deck, reference_profile, include=["LO-006"])
    assert findings_for(result, "LO-006") == []
    # Every shape it declined to measure must be accounted for.
    for entry in result.unchecked:
        if entry.rule_id == "LO-006":
            assert entry.reason, "an unchecked entry must say why"


# --------------------------------------------------------------------------------------
# LO-007 font size band
# --------------------------------------------------------------------------------------


def test_lo007_catches_a_size_outside_the_band(dirty_deck, reference_profile):
    findings = _assert_catches(dirty_deck, reference_profile, "LO-007")
    seeded = [f for f in findings if f.slide_index == SEEDED["LO-007"]]
    assert any("18" in (f.measured or "") for f in seeded), [
        f.measured for f in seeded
    ]


def test_lo007_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    assert_silent_on_clean(_run(clean_deck, reference_profile, "LO-007"), "LO-007")


def test_lo007_does_not_check_a_role_that_was_never_learned(
    clean_deck, reference_profile
):
    """``chart_label`` is in ``not_learned`` for the reference deck, so chart text
    must not be measured against an invented band."""
    assert "chart_label" not in reference_profile.brand.fonts.roles
    assert findings_for(_run(clean_deck, reference_profile, "LO-007"), "LO-007") == []


def test_lo007_respects_an_exact_allowed_set(dirty_deck, reference_profile):
    reference_profile.brand.fonts.roles["body"] = FontRole(exact_pt=[11.0])
    findings = findings_for(_run(dirty_deck, reference_profile, "LO-007"), "LO-007")
    assert findings, "an exact set of 11pt should reject the deck's 10 and 12pt body"


def test_lo007_is_skipped_without_learned_roles(dirty_deck, reference_profile):
    reference_profile.brand.fonts.roles = {}
    result = _run(dirty_deck, reference_profile, "LO-007")
    assert result.findings == []
    assert any(s.rule_id == "LO-007" for s in result.rules_skipped)


# --------------------------------------------------------------------------------------
# LO-008 gutters
# --------------------------------------------------------------------------------------


def test_lo008_catches_uneven_gutters(dirty_deck, reference_profile):
    _assert_catches(dirty_deck, reference_profile, "LO-008")


def test_lo008_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    """The clean deck mixes two, three and four column rows. Each row is evenly
    guttered within itself, which is what the rule measures; the deck has no
    single deck-wide gutter, and the rule must not expect one."""
    assert_silent_on_clean(_run(clean_deck, reference_profile, "LO-008"), "LO-008")


# --------------------------------------------------------------------------------------
# The false-positive guard
# --------------------------------------------------------------------------------------


def test_layout_category_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    """The single most important assertion in this module.

    Every layout rule, enabled ones and LO-006 alike, run against the deck the
    profile describes. Anything reported here is a false positive.
    """
    reference_profile.rules.disabled = []
    reference_profile.rules.enabled = ["LO-006"]
    result = run_rules(clean_deck, reference_profile, include=["LO-*"])
    assert result.findings == [], "; ".join(
        f"slide {f.slide_index} {f.rule_id}: {f.message}" for f in result.findings
    )


def test_layout_rules_exclude_furniture(clean_deck, reference_profile):
    """The logo, page number and confidentiality line sit outside the content
    frame by design, so admitting them would make LO-002 report every slide."""
    furniture = detect_furniture(clean_deck, reference_profile)
    assert furniture.logo_sha1s, "the fixture is supposed to carry a logo"
    assert furniture.page_numbers, "the fixture is supposed to carry page numbers"
    result = run_rules(clean_deck, reference_profile, include=["LO-002", "LO-003"])
    assert result.findings == []


@pytest.mark.parametrize("rule_id", LAYOUT_RULE_IDS)
def test_every_layout_rule_reaches_a_verdict_on_the_dirty_deck(
    rule_id, dirty_deck, reference_profile
):
    """A rule must either run, or say why it did not. Never both silent and absent."""
    reference_profile.rules.disabled = []
    reference_profile.rules.enabled = [rule_id]
    result = run_rules(dirty_deck, reference_profile, include=[rule_id])
    accounted = rule_id in result.rules_run or any(
        s.rule_id == rule_id for s in result.rules_skipped
    )
    assert accounted, f"{rule_id} neither ran nor reported why not"


# --------------------------------------------------------------------------------------
# LO-008's overlap false positive
# --------------------------------------------------------------------------------------


def _deck_of_boxes(path, boxes):
    """A one-slide deck of text boxes at the given ``(left, top, w, h)`` points."""
    from pptx import Presentation
    from pptx.util import Emu, Pt

    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, (left, top, width, height) in enumerate(boxes):
        shape = slide.shapes.add_textbox(Pt(left), Pt(top), Pt(width), Pt(height))
        shape.name = f"Box {index + 1}"
        shape.text_frame.text = f"Box {index + 1}"
    presentation.save(str(path))

    from tieout.model.loader import load_deck

    return load_deck(path)


def test_lo008_does_not_read_a_label_over_a_card_as_its_neighbour(tmp_path, reference_profile):
    """A KPI card and the text box drawn on top of it share a top edge and a
    size, so they pass the sibling test and are read as consecutive members of a
    row. Their "gutter" is then negative, and a real deck produced
    "gutters across a 4-shape row vary by 258.3pt (-20.2pt, 427.2pt, -20.2pt)"
    — a sentence with no meaning for the reader to act on.

    Shapes that overlap are stacked, not set out in a row, and the overlap
    itself is LO-004's to report.
    """
    deck = _deck_of_boxes(
        tmp_path / "stacked.pptx",
        [
            (36.0, 100.0, 400.0, 80.0),   # card
            (56.0, 100.0, 400.0, 80.0),   # its label, drawn over it
            (520.0, 100.0, 400.0, 80.0),  # second card
            (540.0, 100.0, 400.0, 80.0),  # its label
        ],
    )
    assert findings_for(_run(deck, reference_profile, "LO-008"), "LO-008") == []


def test_lo008_still_reports_a_genuinely_uneven_row(tmp_path, reference_profile):
    """The guard above must not have turned the rule off: four same-sized cards
    laid out in a row with one gap wrong is exactly what LO-008 is for."""
    deck = _deck_of_boxes(
        tmp_path / "uneven.pptx",
        [
            (36.0, 100.0, 180.0, 80.0),
            (236.0, 100.0, 180.0, 80.0),   # gutter 20pt
            (436.0, 100.0, 180.0, 80.0),   # gutter 20pt
            (736.0, 100.0, 180.0, 80.0),   # gutter 120pt
        ],
    )
    findings = findings_for(_run(deck, reference_profile, "LO-008"), "LO-008")
    assert findings, "an unevenly guttered row of siblings must still be reported"
    assert "vary by" in findings[0].message


# --------------------------------------------------------------------------------------
# LO-009 bullet indent
# --------------------------------------------------------------------------------------


def _deck_with_bullets(path, indents, *, levels=None):
    """One slide, one list, with each bullet's left indent set explicitly."""
    from pptx import Presentation
    from pptx.util import Emu, Pt

    from tieout.model.loader import load_deck

    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Pt(60), Pt(80), Pt(600), Pt(300))
    box.name = "Body"
    frame = box.text_frame
    for index, indent in enumerate(indents):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"Bullet point number {index + 1} of the list"
        paragraph.level = (levels or [0] * len(indents))[index]
        paragraph._pPr.set("marL", str(int(Pt(indent))))
        paragraph._pPr.set("indent", str(int(Pt(-12))))
        _bullet(paragraph)
    presentation.save(str(path))
    return load_deck(path)


def _bullet(paragraph):
    """Give a paragraph an actual bullet glyph.

    An indented paragraph is not a bulleted one: the model only reads a bullet
    where the file declares one, which is what stops LO-009 reporting an
    indented block quote as a misaligned list.
    """
    from pptx.oxml.ns import qn

    char = paragraph._pPr.makeelement(qn("a:buChar"), {"char": "\u2022"})
    paragraph._pPr.append(char)


def test_lo009_catches_one_bullet_tabbed_out_of_line(tmp_path, reference_profile):
    """A bullet 4pt right of its siblings is a tab somebody pressed, and it is
    visible to a reader at a glance even though no other rule looks at it."""
    deck = _deck_with_bullets(tmp_path / "tabbed.pptx", [27.0, 27.0, 31.0, 27.0])
    findings = findings_for(_run(deck, reference_profile, "LO-009"), "LO-009")

    assert len(findings) == 1
    assert "31pt" in findings[0].message
    assert findings[0].remedy and "27pt" in findings[0].remedy


def test_lo009_is_silent_on_a_list_that_lines_up(tmp_path, reference_profile):
    deck = _deck_with_bullets(tmp_path / "straight.pptx", [27.0, 27.0, 27.0, 27.0])
    assert findings_for(_run(deck, reference_profile, "LO-009"), "LO-009") == []


def test_lo009_does_not_read_a_sub_bullet_as_a_misaligned_one(
    tmp_path, reference_profile
):
    """The level comes from the paragraph, not from the indent. A properly
    demoted sub-bullet is indented on purpose, and reporting it would make the
    rule fire on every nested list in the deck."""
    deck = _deck_with_bullets(
        tmp_path / "nested.pptx",
        [27.0, 54.0, 54.0, 54.0, 27.0, 27.0, 27.0],
        levels=[0, 1, 1, 1, 0, 0, 0],
    )
    assert findings_for(_run(deck, reference_profile, "LO-009"), "LO-009") == []


def test_lo009_needs_enough_bullets_for_disagreement_to_mean_anything(
    tmp_path, reference_profile
):
    """Two bullets at two indents are a two-level list, not a defect."""
    deck = _deck_with_bullets(tmp_path / "two.pptx", [27.0, 31.0])
    assert findings_for(_run(deck, reference_profile, "LO-009"), "LO-009") == []


def test_lo009_tolerates_sub_point_rounding(tmp_path, reference_profile):
    """PowerPoint writes indents in EMU, so a deck that has been resized carries
    fractional differences nobody typed and nobody can see."""
    deck = _deck_with_bullets(tmp_path / "rounded.pptx", [27.0, 27.4, 26.7, 27.1])
    assert findings_for(_run(deck, reference_profile, "LO-009"), "LO-009") == []


def test_lo009_compares_within_one_list_only(tmp_path, reference_profile):
    """A sidebar's bullets and a body's bullets are set to different indents on
    purpose in most house styles. Comparing them reports every deck with a
    sidebar."""
    from pptx import Presentation
    from pptx.util import Emu, Pt

    from tieout.model.loader import load_deck

    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for column, indent in ((60, 27.0), (520, 45.0)):
        box = slide.shapes.add_textbox(Pt(column), Pt(80), Pt(380), Pt(300))
        frame = box.text_frame
        for index in range(4):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = f"Point {index} in the column at {column}"
            paragraph._pPr.set("marL", str(int(Pt(indent))))
            _bullet(paragraph)
    presentation.save(str(tmp_path / "columns.pptx"))

    deck = load_deck(tmp_path / "columns.pptx")
    assert findings_for(_run(deck, reference_profile, "LO-009"), "LO-009") == []


def test_lo009_is_silent_on_the_clean_deck(clean_deck, reference_profile):
    assert_silent_on_clean(_run(clean_deck, reference_profile, "LO-009"), "LO-009")



# --------------------------------------------------------------------------------------
# A geometric finding is as sure as the measurement under it
# --------------------------------------------------------------------------------------


def _deck_of_overlapping_boxes(path, *, filled: bool):
    """Two text boxes that overlap by half, in a face nothing here can measure."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    from tieout.model.loader import load_deck

    presentation = Presentation()
    presentation.slide_width = Emu(960 * 12700)
    presentation.slide_height = Emu(540 * 12700)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, left in enumerate((100.0, 200.0)):
        shape = slide.shapes.add_textbox(Pt(left), Pt(100), Pt(200), Pt(60))
        shape.name = f"Box {index + 1}"
        shape.text_frame.text = f"Box {index + 1} carries a full line of body text"
        for run in shape.text_frame.paragraphs[0].runs:
            run.font.name = "NoSuchFaceEverInstalled"
            run.font.size = Pt(12)
        if filled:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x33)
    presentation.save(str(path))
    return load_deck(path)


def test_an_overlap_between_bounded_text_boxes_is_medium_confidence(
    tmp_path, reference_profile
):
    """Each rectangle is an upper bound on its ink, so two bounds overlapping
    is not two inks overlapping. The finding says so."""
    deck = _deck_of_overlapping_boxes(tmp_path / "bounded.pptx", filled=False)
    findings = _run(deck, reference_profile, "LO-004").findings
    assert findings, "the boxes overlap by half; LO-004 must report it"
    assert findings[0].confidence == "medium"


def test_an_overlap_between_filled_shapes_stays_high_confidence(
    tmp_path, reference_profile
):
    """A fill paints the frame, so the frame is the ink and the overlap is real."""
    deck = _deck_of_overlapping_boxes(tmp_path / "filled.pptx", filled=True)
    findings = _run(deck, reference_profile, "LO-004").findings
    assert findings
    assert findings[0].confidence == "high"
