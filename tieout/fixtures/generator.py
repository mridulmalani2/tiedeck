"""Synthetic reference deck builder.

Serves three purposes: it produces the test fixtures, it produces the deck the
round-trip test learns from and checks against, and it backs the
``tieout scaffold-reference`` command so a new user can see the tool work before
they trust it with a real deck.

Everything is invented. No text, figure, company name or codename here comes from
any real filing, bank or client.

Two variants, per section 11:

* ``clean`` -- fully self-consistent. This is the golden deck and the
  false-positive guard. ``tieout check`` against a profile learned from it must
  return nothing at all.
* ``dirty`` -- the clean deck plus one seeded violation per rule, each tagged in
  the spec with its rule id.

Four defects cannot share a deck with the others and get their own single-defect
variants: a changed slide size invalidates every geometric seed, and the three
package-level hygiene defects are applied by rewriting the archive after
python-pptx has closed it.
"""

from __future__ import annotations

import io
import re
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Final

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from tieout.fixtures.spec import (
    BODY_BOTTOM_PT,
    BODY_TOP_PT,
    COLUMNS_2_PT,
    COLUMNS_3_PT,
    COLUMNS_4_PT,
    CONTENT_LEFT_PT,
    CONTENT_WIDTH_PT,
    FOOTNOTE_RULE_TOP_PT,
    FOOTNOTE_TOP_PT,
    TITLE_HEIGHT_PT,
    TITLE_RULE_TOP_PT,
    TITLE_TOP_PT,
    BrandSpec,
    FontSpec,
    ReferenceSpec,
    SlideSpec,
    default_spec,
)
from tieout.model.color import parse_hex
from tieout.model.inherit import NS

A: Final[str] = NS["a"]
P: Final[str] = NS["p"]

#: "No Style, No Grid" -- the only built-in table style that contributes no
#: colours of its own, so the fixture's palette stays exactly what the spec says.
_PLAIN_TABLE_STYLE: Final[str] = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

#: Layout indices in the stock python-pptx template that the generator repurposes.
#: Where a recap table sits: bottom right of a content slide, clear of the two
#: body columns and above the footnote rule. Every seeded tie-out defect that
#: needs a table of its own goes here, so none of them collides with the body
#: text the layout rules are measured against.
_RECAP_LEFT_PT: Final[int] = 492
_RECAP_TOP_PT: Final[int] = 384
_RECAP_WIDTH_PT: Final[int] = 432

_LAYOUT_TITLE: Final[int] = 0
_LAYOUT_SECTION: Final[int] = 2
_LAYOUT_TITLE_ONLY: Final[int] = 5
_LAYOUT_BLANK: Final[int] = 6

_LAYOUT_PLAN: Final[tuple[tuple[int, str, str], ...]] = (
    (_LAYOUT_TITLE, "Title Slide", "title"),
    (_LAYOUT_SECTION, "Section Divider", "secHead"),
    (_LAYOUT_TITLE_ONLY, "Content", "titleOnly"),
    (_LAYOUT_BLANK, "Blank", "blank"),
)

_DISCLAIMER_PARAGRAPHS: Final[tuple[str, ...]] = (
    "These materials have been prepared by Ashcombe Partners solely for the "
    "information of the Board of Directors of the company referred to herein and "
    "may not be relied upon by any other person for any purpose. They are "
    "incomplete without the accompanying oral presentation and must be considered "
    "only in conjunction with it.",
    "The information contained in these materials has been obtained from the "
    "company and from sources believed to be reliable. Ashcombe Partners has not "
    "independently verified that information and makes no representation or "
    "warranty, express or implied, as to its accuracy, completeness or "
    "reasonableness.",
    "Any estimates, projections or forecasts contained in these materials involve "
    "significant elements of subjective judgement and analysis. They are not "
    "necessarily indicative of future results and no representation is made that "
    "any such estimate, projection or forecast will be realised.",
    "Ashcombe Partners does not provide accounting, tax or legal advice. Recipients "
    "should consult their own advisers in respect of those matters. These materials "
    "do not constitute an offer to sell or a solicitation of an offer to buy any "
    "security and may not be reproduced or circulated without prior written consent.",
)

_AGENDA_NUMERALS: Final[tuple[str, ...]] = ("I", "II", "III", "IV")


class GeneratorError(RuntimeError):
    """Raised when the fixture cannot be built as specified."""


@dataclass(frozen=True)
class BuildResult:
    """What a build produced, so tests can assert against the same paths."""

    clean: Path
    dirty: Path
    variants: dict[str, Path]
    spec_path: Path


# --------------------------------------------------------------------------------------
# Public entry points
# --------------------------------------------------------------------------------------


def build_clean(path: str | Path, spec: ReferenceSpec | None = None) -> Path:
    """Build the self-consistent golden deck."""
    return _build(Path(path), spec or default_spec(), seed_defects=False, variant=None)


def build_dirty(path: str | Path, spec: ReferenceSpec | None = None) -> Path:
    """Build the golden deck plus every co-existable seeded defect."""
    return _build(Path(path), spec or default_spec(), seed_defects=True, variant=None)


def build_variant(
    path: str | Path, variant: str, spec: ReferenceSpec | None = None
) -> Path:
    """Build a single-defect deck for a defect that cannot share a deck."""
    return _build(Path(path), spec or default_spec(), seed_defects=False, variant=variant)


def build_all(directory: str | Path, spec: ReferenceSpec | None = None) -> BuildResult:
    """Build every fixture into ``directory`` and write the spec beside them."""
    spec = spec or default_spec()
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)

    clean = build_clean(directory / "reference_clean.pptx", spec)
    dirty = build_dirty(directory / "reference_dirty.pptx", spec)

    variants: dict[str, Path] = {}
    for name in sorted({d.variant for d in spec.defects if d.variant}):
        variants[name] = build_variant(directory / f"reference_{name}.pptx", name, spec)

    spec_path = directory / "reference_spec.yaml"
    spec_path.write_text(spec.to_yaml(), encoding="utf-8")
    return BuildResult(clean=clean, dirty=dirty, variants=variants, spec_path=spec_path)


# --------------------------------------------------------------------------------------
# Build
# --------------------------------------------------------------------------------------


def _build(
    path: Path, spec: ReferenceSpec, *, seed_defects: bool, variant: str | None
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()

    width = spec.width_pt if variant != "wrong_slide_size" else 720
    height = spec.height_pt
    presentation.slide_width = Emu(int(width * 12700))
    presentation.slide_height = Emu(int(height * 12700))

    _restyle_theme(presentation, spec.brand)
    _restyle_master(presentation, spec.brand)
    _restyle_layouts(presentation, spec.brand, width)

    logo_bytes = _logo_png(spec.brand)
    low_res_bytes = _logo_png(spec.brand, pixels=(40, 14))

    defects = _defect_index(spec) if seed_defects else {}
    builder = _DeckBuilder(
        presentation=presentation,
        spec=spec,
        width_pt=float(width),
        height_pt=float(height),
        logo_bytes=logo_bytes,
        low_res_logo_bytes=low_res_bytes,
        defects=defects,
    )
    for slide_spec in spec.slides:
        builder.add_slide(slide_spec)

    presentation.save(str(path))
    _rewrite_package(path, spec, defects=defects, variant=variant)
    return path


def _defect_index(spec: ReferenceSpec) -> dict[str, int | None]:
    """Rule id -> slide index, for defects that live on a slide."""
    return {d.rule_id: d.slide_index for d in spec.defects if d.variant is None}


# --------------------------------------------------------------------------------------
# Theme, master and layout restyling
# --------------------------------------------------------------------------------------


def _restyle_theme(presentation: Any, brand: BrandSpec) -> None:
    """Point the theme's font scheme and colour scheme at the house style.

    Without this, any property TieOut resolves through the theme rather than
    through an explicit run lands on Calibri and an Office accent colour, and the
    learner would observe those as part of the client's brand.
    """
    for master in presentation.slide_masters:
        try:
            theme_part = master.part.part_related_by(
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
            )
        except KeyError:  # pragma: no cover - stock template always has a theme
            continue
        root = etree.fromstring(theme_part.blob)

        for which, typeface in (
            ("majorFont", brand.primary_font),
            ("minorFont", brand.primary_font),
        ):
            node = root.find(f".//a:themeElements/a:fontScheme/a:{which}/a:latin", NS)
            if node is not None:
                node.set("typeface", typeface)

        scheme_colours = {
            "dk1": brand.ink,
            "lt1": brand.paper,
            "dk2": brand.house_navy,
            "lt2": brand.rule_grey,
            "accent1": brand.house_navy,
            "accent2": brand.rule_grey,
            "accent3": brand.accent_red,
            "accent4": brand.house_navy,
            "accent5": brand.rule_grey,
            "accent6": brand.accent_red,
            "hlink": brand.house_navy,
            "folHlink": brand.house_navy,
        }
        for slot, hex_value in scheme_colours.items():
            slot_node = root.find(f".//a:themeElements/a:clrScheme/a:{slot}", NS)
            if slot_node is None:
                continue
            for child in list(slot_node):
                slot_node.remove(child)
            srgb = etree.SubElement(slot_node, f"{{{A}}}srgbClr")
            srgb.set("val", hex_value.lstrip("#").upper())

        theme_part._blob = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )


def _restyle_master(presentation: Any, brand: BrandSpec) -> None:
    """Rewrite the master's text styles to the house style.

    This makes the inheritance chain coherent: a run with no explicit properties
    resolves to the same values as a run with them, so a resolution bug shows up
    as a mismatch rather than being masked by a fallback that happens to agree.
    """
    for master in presentation.slide_masters:
        styles = master.element.find(f"{{{P}}}txStyles")
        if styles is None:
            continue
        plan = (
            ("titleStyle", brand.title_pt, brand.house_navy, True),
            ("bodyStyle", brand.body_sizes_pt[1], brand.ink, False),
            ("otherStyle", brand.body_sizes_pt[1], brand.ink, False),
        )
        for tag, size_pt, colour_hex, bold in plan:
            style = styles.find(f"{{{P}}}{tag}")
            if style is None:
                continue
            for level in range(1, 10):
                lvl = style.find(f"{{{A}}}lvl{level}pPr")
                if lvl is None:
                    continue
                _write_def_rpr(lvl, brand.primary_font, size_pt, colour_hex, bold)


def _write_def_rpr(
    lvl_ppr: etree._Element,
    typeface: str,
    size_pt: float,
    colour_hex: str,
    bold: bool,
) -> None:
    existing = lvl_ppr.find(f"{{{A}}}defRPr")
    if existing is not None:
        lvl_ppr.remove(existing)
    def_rpr = etree.SubElement(lvl_ppr, f"{{{A}}}defRPr")
    def_rpr.set("sz", str(round(size_pt * 100)))
    def_rpr.set("b", "1" if bold else "0")
    solid = etree.SubElement(def_rpr, f"{{{A}}}solidFill")
    srgb = etree.SubElement(solid, f"{{{A}}}srgbClr")
    srgb.set("val", colour_hex.lstrip("#").upper())
    latin = etree.SubElement(def_rpr, f"{{{A}}}latin")
    latin.set("typeface", typeface)


def _restyle_layouts(presentation: Any, brand: BrandSpec, width_pt: float) -> None:
    """Rename and retype the layouts, and move their title placeholders.

    The stock template is 4:3, so its layout placeholders sit at 4:3 positions.
    Leaving them there would make BR-008 report drift on every slide in a 16:9
    deck, because the slide title and its layout placeholder would genuinely
    disagree.
    """
    layouts = presentation.slide_layouts
    for index, name, layout_type in _LAYOUT_PLAN:
        if index >= len(layouts):  # pragma: no cover - stock template has 11
            continue
        layout = layouts[index]
        c_sld = layout.element.find(f"{{{P}}}cSld")
        if c_sld is not None:
            c_sld.set("name", name)
        layout.element.set("type", layout_type)

        content_width = width_pt - 2 * CONTENT_LEFT_PT
        for placeholder in layout.placeholders:
            ph_type = str(placeholder.element.find(".//p:nvPr/p:ph", NS).get("type") or "body")
            if ph_type in ("title", "ctrTitle"):
                if layout_type == "title":
                    _set_geometry(
                        placeholder, CONTENT_LEFT_PT, 216, content_width, 60
                    )
                elif layout_type == "secHead":
                    _set_geometry(
                        placeholder, CONTENT_LEFT_PT, 300, content_width, 36
                    )
                else:
                    _set_geometry(
                        placeholder,
                        CONTENT_LEFT_PT,
                        TITLE_TOP_PT,
                        content_width,
                        TITLE_HEIGHT_PT,
                    )
            elif ph_type == "subTitle":
                _set_geometry(placeholder, CONTENT_LEFT_PT, 288, content_width, 24)
            elif ph_type == "body" and layout_type == "secHead":
                _set_geometry(placeholder, CONTENT_LEFT_PT, 348, content_width, 24)


def _set_geometry(
    shape: Any, left_pt: float, top_pt: float, width_pt: float, height_pt: float
) -> None:
    shape.left = Emu(int(left_pt * 12700))
    shape.top = Emu(int(top_pt * 12700))
    shape.width = Emu(int(width_pt * 12700))
    shape.height = Emu(int(height_pt * 12700))


# --------------------------------------------------------------------------------------
# Logo
# --------------------------------------------------------------------------------------


def _logo_png(brand: BrandSpec, pixels: tuple[int, int] | None = None) -> bytes:
    """A deterministic wordmark-ish PNG, sized so effective DPI clears 150.

    Drawn rather than shipped as a binary, because section 12 forbids committing
    binary fixtures.
    """
    width, height = pixels or brand.logo_pixels
    image = Image.new("RGB", (width, height), parse_hex(brand.paper).as_tuple())
    draw = ImageDraw.Draw(image)
    navy = parse_hex(brand.house_navy).as_tuple()
    bar = max(2, height // 6)
    draw.rectangle([0, 0, width - 1, bar], fill=navy)
    draw.rectangle(
        [0, height - bar - 1, width - 1, height - 1], fill=navy
    )
    draw.rectangle(
        [bar * 2, bar * 2, width // 3, height - bar * 2 - 1], fill=navy
    )
    draw.rectangle(
        [width // 2, bar * 2, width - bar * 2 - 1, height // 2], fill=navy
    )
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


# --------------------------------------------------------------------------------------
# Deck builder
# --------------------------------------------------------------------------------------


class _DeckBuilder:
    """Builds slides onto a presentation from :class:`SlideSpec` entries."""

    def __init__(
        self,
        *,
        presentation: Any,
        spec: ReferenceSpec,
        width_pt: float,
        height_pt: float,
        logo_bytes: bytes,
        low_res_logo_bytes: bytes,
        defects: dict[str, int | None],
    ) -> None:
        self.presentation = presentation
        self.spec = spec
        self.brand = spec.brand
        self.width_pt = width_pt
        self.height_pt = height_pt
        self.logo_bytes = logo_bytes
        self.low_res_logo_bytes = low_res_logo_bytes
        self.defects = defects
        self.page_numbers_emitted: list[int] = []

    # -- helpers -----------------------------------------------------------------

    def defect_on(self, rule_id: str, slide_index: int) -> bool:
        return self.defects.get(rule_id) == slide_index

    def _layout(self, kind: str) -> Any:
        index = {
            "title": _LAYOUT_TITLE,
            "divider": _LAYOUT_SECTION,
        }.get(kind, _LAYOUT_TITLE_ONLY)
        return self.presentation.slide_layouts[index]

    def add_slide(self, slide_spec: SlideSpec) -> Any:
        slide = self.presentation.slides.add_slide(self._layout(slide_spec.kind))
        builders = {
            "title": self._build_title,
            "agenda": self._build_agenda,
            "divider": self._build_divider,
            "content": self._build_content,
            "table": self._build_table,
            "chart": self._build_chart,
            "disclaimer": self._build_disclaimer,
        }
        builder = builders.get(slide_spec.kind)
        if builder is None:
            raise GeneratorError(f"unknown slide kind: {slide_spec.kind}")
        builder(slide, slide_spec)
        self._add_furniture(slide, slide_spec)
        if self.defect_on("HY-002", slide_spec.index):
            slide.notes_slide.notes_text_frame.text = (
                "Check the depot capacity figure with the client before Thursday."
            )
        if self.defect_on("HY-003", slide_spec.index):
            slide.element.set("show", "0")
        return slide

    # -- furniture ---------------------------------------------------------------

    def _add_furniture(self, slide: Any, slide_spec: SlideSpec) -> None:
        """Logo, page number and confidentiality line."""
        if slide_spec.has_logo and not self.defect_on("BR-001", slide_spec.index):
            box = (
                self.brand.logo_divider_box_pt
                if slide_spec.logo_centred
                else self.brand.logo_box_pt
            )
            left, top, width, height = box
            if self.defect_on("BR-002", slide_spec.index):
                left -= 18
            if self.defect_on("BR-003", slide_spec.index):
                width = int(width * 1.35)
            payload = (
                self.low_res_logo_bytes
                if self.defect_on("HY-008", slide_spec.index)
                else self.logo_bytes
            )
            picture = slide.shapes.add_picture(
                io.BytesIO(payload), Pt(left), Pt(top), Pt(width), Pt(height)
            )
            picture.name = "Logo"

        if slide_spec.has_page_number:
            left, top, width, height = self.brand.page_number_box_pt
            number = str(slide_spec.index)
            if self.defect_on("BR-006", slide_spec.index):
                number = f"Page {slide_spec.index} of 26"
            if self.defect_on("BR-007", slide_spec.index):
                number = "3"
            box = self._text_box(
                slide,
                left,
                top,
                width,
                height,
                "Page number",
                align=PP_ALIGN.LEFT,
            )
            self._write_paragraph(
                box.text_frame.paragraphs[0], number, self.brand.footnote_font()
            )
            self.page_numbers_emitted.append(slide_spec.index)

        if slide_spec.has_confidentiality and not self.defect_on(
            "BR-010", slide_spec.index
        ):
            left, top, width, height = self.brand.confidentiality_box_pt
            box = self._text_box(
                slide, left, top, width, height, "Confidentiality", align=PP_ALIGN.RIGHT
            )
            self._write_paragraph(
                box.text_frame.paragraphs[0],
                self.brand.confidentiality_text,
                self.brand.footnote_font(),
            )

    # -- slide kinds -------------------------------------------------------------

    def _build_title(self, slide: Any, slide_spec: SlideSpec) -> None:
        title = slide.shapes.title
        if title is not None:
            self._fill_placeholder(
                title, slide_spec.title, self.brand.title_font(), align=PP_ALIGN.LEFT
            )
        subtitle = self._placeholder_of(slide, "subTitle")
        if subtitle is not None:
            self._fill_placeholder(
                subtitle,
                slide_spec.subtitle,
                self.brand.subtitle_font(),
                align=PP_ALIGN.LEFT,
            )
        else:  # pragma: no cover - stock title layout always has a subtitle
            box = self._text_box(
                slide, CONTENT_LEFT_PT, 288, CONTENT_WIDTH_PT, 24, "Subtitle"
            )
            self._write_paragraph(
                box.text_frame.paragraphs[0],
                slide_spec.subtitle,
                self.brand.subtitle_font(),
            )

        self._rule(slide, CONTENT_LEFT_PT, 276, CONTENT_WIDTH_PT, self.brand.house_navy)
        detail = self._text_box(
            slide, CONTENT_LEFT_PT, 324, CONTENT_WIDTH_PT, 36, "Advisor mark"
        )
        frame = detail.text_frame
        self._write_paragraph(
            frame.paragraphs[0], self.spec.advisor_mark, self.brand.body_font(12.0)
        )
        self._write_paragraph(
            frame.add_paragraph(), self.spec.deck_date, self.brand.body_font(10.0)
        )

    def _build_agenda(self, slide: Any, slide_spec: SlideSpec) -> None:
        self._title_block(slide, slide_spec)
        sections: list[str] = list(slide_spec.payload.get("sections", []))
        top = BODY_TOP_PT
        for numeral, entry in zip(_AGENDA_NUMERALS, sections, strict=False):
            heading = entry.split(". ", 1)[-1]
            box = self._text_box(
                slide, CONTENT_LEFT_PT, top, CONTENT_WIDTH_PT, 24, f"Agenda {numeral}"
            )
            paragraph = box.text_frame.paragraphs[0]
            self._write_paragraph(
                paragraph, f"{numeral}.", self.brand.body_font(14.0), no_bullet=True
            )
            run = paragraph.add_run()
            run.text = f"   {heading}"
            self._apply_font(run.font, self.brand.body_font(14.0))
            top += 48

    def _build_divider(self, slide: Any, slide_spec: SlideSpec) -> None:
        title = slide.shapes.title
        if title is not None:
            self._fill_placeholder(
                title, slide_spec.title, self.brand.title_font(), align=PP_ALIGN.CENTER
            )
        body = self._placeholder_of(slide, "body")
        if body is not None:
            self._fill_placeholder(
                body,
                slide_spec.subtitle,
                self.brand.subtitle_font(),
                align=PP_ALIGN.CENTER,
            )
        self._rule(slide, 420, 336, 120, self.brand.house_navy)

    def _build_content(self, slide: Any, slide_spec: SlideSpec) -> None:
        self._title_block(slide, slide_spec)
        column_count = int(slide_spec.payload.get("columns", 2))
        columns = {2: COLUMNS_2_PT, 3: COLUMNS_3_PT, 4: COLUMNS_4_PT}[column_count]
        bullet_groups: list[list[str]] = list(slide_spec.payload.get("bullets", []))

        if self.defect_on("LO-008", slide_spec.index):
            columns = tuple(
                (left + (9 if position == 2 else 0), width)
                for position, (left, width) in enumerate(columns)
            )
        if self.defect_on("LO-003", slide_spec.index):
            columns = tuple(
                (left + (3 if position == 1 else 0), width)
                for position, (left, width) in enumerate(columns)
            )

        for position, (bullets, (left, width)) in enumerate(
            zip(bullet_groups, columns, strict=False)
        ):
            name = f"Column {position + 1}"
            height = BODY_BOTTOM_PT - BODY_TOP_PT
            box = self._text_box(
                slide, left, BODY_TOP_PT, width, height, name
            )
            frame = box.text_frame
            frame.word_wrap = True
            size = self.brand.body_sizes_pt[1]
            if column_count == 4:
                # A four-column row is tighter, so the house style drops a point.
                # This also gives the body role four distinct observed sizes, which
                # is what makes the learner emit a band rather than an exact set.
                size = self.brand.body_sizes_pt[0]
            elif column_count == 3 and position == 0:
                size = self.brand.body_sizes_pt[2]
            if self.defect_on("LO-007", slide_spec.index) and position == 0:
                size = 18.0
            for order, text in enumerate(bullets):
                paragraph = (
                    frame.paragraphs[0] if order == 0 else frame.add_paragraph()
                )
                # Curl the quotes first, then seed the defects. Seeding a straight
                # apostrophe before the quote pass would have the generator curl
                # its own defect away, which is exactly what happened.
                body_text = self._curly(text)
                emphasis = column_count >= 3 and order == 0
                if self.defect_on("TY-001", slide_spec.index) and order == 0:
                    body_text = body_text + " against the company's position"
                if self.defect_on("TY-002", slide_spec.index) and order == 0:
                    body_text = body_text.replace(" ", "  ", 1) + " ."
                if self.defect_on("TY-003", slide_spec.index) and order == 0:
                    body_text = body_text + "."
                if self.defect_on("TY-005", slide_spec.index) and order == 0:
                    body_text = body_text.replace(
                        "broad process", "Ashcombe partners process"
                    )
                    if "Ashcombe partners" not in body_text:
                        body_text = f"{body_text} per Ashcombe partners"
                if self.defect_on("TY-009", slide_spec.index) and order == 0:
                    body_text = f"{body_text}, subbject to confirmation"
                if self.defect_on("HY-001", slide_spec.index) and order == 0:
                    body_text = f"{body_text} [TBD]"
                font = FontSpec(
                    self.brand.primary_font,
                    size,
                    self.brand.house_navy if emphasis else self.brand.ink,
                    bold=emphasis,
                )
                self._write_paragraph(
                    paragraph,
                    body_text,
                    font,
                    bullet_char=None if emphasis else "\u2013",
                )

        if self.defect_on("LO-004", slide_spec.index):
            overlap = self._text_box(
                slide, CONTENT_LEFT_PT + 24, BODY_TOP_PT + 24, 300, 48, "Callout"
            )
            self._write_paragraph(
                overlap.text_frame.paragraphs[0],
                "Overlapping callout box",
                self.brand.body_font(11.0),
            )
        if self.defect_on("LO-001", slide_spec.index):
            spill = self._text_box(
                slide, self.width_pt - 60, BODY_TOP_PT, 180, 24, "Off canvas"
            )
            self._write_paragraph(
                spill.text_frame.paragraphs[0],
                "This box runs off the canvas",
                self.brand.body_font(11.0),
            )
        if self.defect_on("LO-002", slide_spec.index):
            intruder = self._text_box(slide, 8, BODY_TOP_PT, 180, 24, "Margin intruder")
            self._write_paragraph(
                intruder.text_frame.paragraphs[0],
                "Inside the safe margin",
                self.brand.body_font(11.0),
            )
        if self.defect_on("LO-006", slide_spec.index):
            cramped = self._text_box(slide, CONTENT_LEFT_PT, 384, 120, 24, "Cramped")
            frame = cramped.text_frame
            frame.word_wrap = True
            self._write_paragraph(
                frame.paragraphs[0],
                "This sentence is far too long to fit inside a box this small and "
                "will overflow it comprehensively",
                self.brand.figure_cell_font(10.0),
            )
        if self.defect_on("BR-004", slide_spec.index):
            swatch = self._rect(slide, 720, 384, 168, 48, "#2E8B7A", "Off palette box")
            swatch.name = "Off palette box"
        if self.defect_on("BR-005", slide_spec.index):
            box = self._text_box(slide, 720, 384, 168, 24, "Unapproved font")
            self._write_paragraph(
                box.text_frame.paragraphs[0],
                "Set in Comic Sans MS",
                FontSpec("Comic Sans MS", 11.0, self.brand.ink),
            )
        if self.defect_on("HY-006", slide_spec.index):
            self._add_empty_placeholder(slide)
        if self.defect_on("HY-009", slide_spec.index):
            box = self._text_box(slide, 720, 420, 168, 24, "Exotic font")
            self._write_paragraph(
                box.text_frame.paragraphs[0],
                "Set in Bodoni Sixtysix",
                FontSpec("Bodoni Sixtysix", 11.0, self.brand.ink),
            )

        if self.defect_on("CO-001", slide_spec.index):
            self._recap_table(slide, "351", "375")
        if self.defect_on("CO-002", slide_spec.index):
            self._recap_table(slide, "1,908", "1,908,000", metric="Revenue")

        # -- the derived checks ----------------------------------------------
        #
        # Every one of these is placed against a period or a label that the
        # projections table on slide 6 does not carry. That is not fussiness:
        # a wrong margin stated for 2025A is *also* a CO-001 contradiction with
        # slide 6's own margin column, and a bridge whose bars are labelled
        # "EBITDA" restates EBITDA levels the projections table already gives.
        # Either would make two rules fire on one seed, and the suite's per-rule
        # assertion -- catches its own, reports nothing else -- would stop
        # meaning anything.
        if self.defect_on("CO-004", slide_spec.index):
            self._grid_table(
                slide,
                ("Fiscal year", "Revenue", "EBITDA", "Margin"),
                # 580 / 3,050 is 19.0%, not 24.0%.
                (("2028E", "3,050", "580", "24.0%"),),
            )
        if self.defect_on("CO-006", slide_spec.index):
            self._grid_table(
                slide,
                ("Metric", "Value"),
                # 2,760 from 1,562 over three years is 20.9%, not 30.0%.
                #
                # Struck from 2024A rather than 2023A deliberately. TY-006 seeds
                # a stray decimal into the 2023A revenue cell, so on the dirty
                # deck that year is stated two ways -- 1,284.5 in the table and
                # 1,284 in the chart -- and CO-006 refuses to compute from an
                # ambiguous endpoint, which is the behaviour it should have and
                # would have made this seed silent.
                (("Revenue CAGR 2024A-2027E", "30.0%"),),
            )
        if self.defect_on("CO-007", slide_spec.index):
            self._grid_table(
                slide,
                # Headed "movement", not "EBITDA": the bars are deltas, and a
                # column headed EBITDA would have every step read as a level
                # and compared with the projections table.
                ("Step", "EBITDA movement"),
                (
                    ("Opening", "263"),
                    ("Volume", "45"),
                    ("Price", "30"),
                    ("Cost", "13"),
                    # 263 + 45 + 30 + 13 is 351, not 400.
                    ("Closing", "400"),
                ),
            )
        if self.defect_on("CO-009", slide_spec.index):
            # Restates one comparable's enterprise value correctly, under a
            # footnote dated ten weeks earlier than slide 12's. The figures
            # agree; only the date they are true of does not.
            self._grid_table(
                slide,
                ("Company", "Enterprise value"),
                (("Calderwood Logistics", "8,420"),),
            )
        if self.defect_on("CO-008", slide_spec.index):
            # The same figure slide 6 states as 1,908 in millions, told in
            # billions. Arithmetically identical, which is the point: CO-001
            # must stay silent and CO-008 must not.
            self._unit_caption(slide, "Figures in US$ billions")
            self._grid_table(
                slide, ("Fiscal year", "Revenue"), (("2025A", "1.908"),)
            )

        self._footnote(slide, slide_spec)

    def _recap_table(
        self, slide: Any, correct: str, seeded: str, metric: str = "EBITDA"
    ) -> None:
        """A small recap table restating a figure from the projections table.

        Decks restate key figures on summary slides constantly, which is exactly
        why a figure can end up disagreeing with itself. The row and column
        labels match slide 6's table so the consistency rules have a genuine
        pair to compare, and only the value differs.
        """
        frame = slide.shapes.add_table(
            2, 2, Pt(_RECAP_LEFT_PT), Pt(_RECAP_TOP_PT), Pt(_RECAP_WIDTH_PT), Pt(48)
        )
        frame.name = "Recap table"
        table = frame.table
        self._plain_table_style(table)
        table.columns[0].width = Emu(216 * 12700)
        table.columns[1].width = Emu(216 * 12700)
        for index in range(2):
            table.rows[index].height = Emu(24 * 12700)

        for column, heading in enumerate(("Fiscal year", metric)):
            cell = table.cell(0, column)
            self._fill_cell(cell, self.brand.house_navy)
            self._cell_text(cell, heading, self.brand.header_cell_font())

        for column, value in enumerate(("2025A", seeded)):
            cell = table.cell(1, column)
            self._fill_cell(cell, self.brand.paper)
            font = (
                self.brand.body_font(self.brand.table_sizes_pt[1])
                if column == 0
                else self.brand.figure_cell_font()
            )
            self._cell_text(cell, value, font)
        del correct

    def _unit_caption(self, slide: Any, text: str) -> None:
        """A scale stated directly above the recap table, overriding the slide's.

        Placed to overlap the table horizontally and to sit just above it,
        because that is the only arrangement :func:`tieout.figures._inherited_unit`
        will take a caption from in preference to the slide's own footnote.
        """
        box = self._text_box(slide, _RECAP_LEFT_PT, _RECAP_TOP_PT - 16, _RECAP_WIDTH_PT,
                             12, "Recap units")
        self._write_paragraph(
            box.text_frame.paragraphs[0], text, self.brand.footnote_font()
        )

    def _grid_table(
        self, slide: Any, headings: tuple[str, ...], rows: tuple[tuple[str, ...], ...]
    ) -> None:
        """A small table in the recap slot, of arbitrary shape.

        :meth:`_recap_table` restates one figure and is fixed at two by two.
        The derived checks need three and four columns, and a bridge needs six
        rows, so the general form lives here and the shape comes from the seed.
        """
        columns = len(headings)
        height = 18 * (len(rows) + 1)
        frame = slide.shapes.add_table(
            len(rows) + 1,
            columns,
            Pt(_RECAP_LEFT_PT),
            Pt(_RECAP_TOP_PT),
            Pt(_RECAP_WIDTH_PT),
            Pt(height),
        )
        frame.name = "Recap table"
        table = frame.table
        self._plain_table_style(table)
        first = _RECAP_WIDTH_PT // 2 if columns == 2 else _RECAP_WIDTH_PT // (columns + 1)
        rest = (_RECAP_WIDTH_PT - first) // max(1, columns - 1)
        for index in range(columns):
            table.columns[index].width = Emu((first if index == 0 else rest) * 12700)
        for index in range(len(rows) + 1):
            table.rows[index].height = Emu(18 * 12700)

        for column, heading in enumerate(headings):
            cell = table.cell(0, column)
            self._fill_cell(cell, self.brand.house_navy)
            self._cell_text(cell, heading, self.brand.header_cell_font())
        for row_index, row in enumerate(rows, start=1):
            for column, value in enumerate(row):
                cell = table.cell(row_index, column)
                self._fill_cell(cell, self.brand.paper)
                font = (
                    self.brand.body_font(self.brand.table_sizes_pt[1])
                    if column == 0
                    else self.brand.figure_cell_font()
                )
                self._cell_text(cell, value, font)

    def _build_table(self, slide: Any, slide_spec: SlideSpec) -> None:
        self._title_block(slide, slide_spec)
        columns: list[str] = list(slide_spec.payload.get("columns", []))
        rows: list[list[str]] = [list(r) for r in slide_spec.payload.get("rows", [])]

        if self.defect_on("CO-003", slide_spec.index) and rows:
            for row in rows:
                if row and row[0].lower().startswith("total"):
                    row[1] = "9,900"
                    break
        if self.defect_on("TY-006", slide_spec.index) and rows:
            rows[0][1] = "1,284.5"
        if self.defect_on("TY-007", slide_spec.index) and rows:
            rows[0][1] = "USD 8,420"
        if self.defect_on("TY-008", slide_spec.index) and rows:
            rows[0][0] = "As at 09/14/2026"
        if self.defect_on("CO-005", slide_spec.index) and len(rows) > 2:
            # 4,935 over 602 is 8.2x, not 11.5x. Seeded on the third row rather
            # than the first because TY-007 already rewrites the first, and two
            # seeds in one cell tell you nothing about either rule.
            rows[2][3] = "11.5x"

        row_count = len(rows) + 1
        table_height = 24 + 24 * len(rows)
        frame = slide.shapes.add_table(
            row_count,
            len(columns),
            Pt(CONTENT_LEFT_PT),
            Pt(BODY_TOP_PT),
            Pt(CONTENT_WIDTH_PT),
            Pt(table_height),
        )
        frame.name = "Financial table"
        table = frame.table
        self._plain_table_style(table)

        first_width = 240
        rest = (CONTENT_WIDTH_PT - first_width) // max(1, len(columns) - 1)
        for index in range(len(columns)):
            table.columns[index].width = Emu(
                int((first_width if index == 0 else rest) * 12700)
            )
        table.rows[0].height = Emu(24 * 12700)
        for index in range(1, row_count):
            table.rows[index].height = Emu(24 * 12700)

        for column_index, heading in enumerate(columns):
            cell = table.cell(0, column_index)
            self._fill_cell(cell, self.brand.house_navy)
            self._cell_text(
                cell,
                heading,
                self.brand.header_cell_font(),
                align=PP_ALIGN.LEFT if column_index == 0 else PP_ALIGN.RIGHT,
            )

        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row):
                cell = table.cell(row_index, column_index)
                self._fill_cell(cell, self.brand.paper)
                is_label = column_index == 0
                font = (
                    self.brand.body_font(self.brand.table_sizes_pt[1])
                    if is_label
                    else self.brand.figure_cell_font()
                )
                self._cell_text(
                    cell,
                    self._curly(value),
                    font,
                    align=PP_ALIGN.LEFT if is_label else PP_ALIGN.RIGHT,
                )

        self._footnote(slide, slide_spec)

    def _build_chart(self, slide: Any, slide_spec: SlideSpec) -> None:
        self._title_block(slide, slide_spec)
        data: Any = CategoryChartData()  # type: ignore[no-untyped-call]
        data.categories = list(slide_spec.payload.get("categories", []))
        for name, values in slide_spec.payload.get("series", []):
            data.add_series(name, tuple(values))

        graphic_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Pt(CONTENT_LEFT_PT),
            Pt(BODY_TOP_PT),
            Pt(CONTENT_WIDTH_PT),
            Pt(BODY_BOTTOM_PT - BODY_TOP_PT),
            data,
        )
        graphic_frame.name = "Chart"
        chart = graphic_frame.chart
        chart.has_title = False
        label_font = self.brand.chart_label_font()
        chart.font.name = label_font.name
        chart.font.size = Pt(label_font.size_pt)
        chart.font.color.rgb = _rgb(label_font.color_hex)
        chart.has_legend = len(chart.plots[0].series) > 1
        if chart.has_legend:
            chart.legend.include_in_layout = False

        palette = (self.brand.house_navy, self.brand.rule_grey)
        for index, series in enumerate(chart.plots[0].series):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = _rgb(palette[index % len(palette)])

        self._footnote(slide, slide_spec)

    def _build_disclaimer(self, slide: Any, slide_spec: SlideSpec) -> None:
        self._title_block(slide, slide_spec)
        count = int(slide_spec.payload.get("paragraph_count", 4))
        box = self._text_box(
            slide,
            CONTENT_LEFT_PT,
            BODY_TOP_PT,
            CONTENT_WIDTH_PT,
            BODY_BOTTOM_PT - BODY_TOP_PT,
            "Disclaimer body",
        )
        frame = box.text_frame
        frame.word_wrap = True
        for order, text in enumerate(_DISCLAIMER_PARAGRAPHS[:count]):
            paragraph = frame.paragraphs[0] if order == 0 else frame.add_paragraph()
            self._write_paragraph(
                paragraph,
                self._curly(text),
                self.brand.footnote_font(),
                no_bullet=True,
            )

    # -- shared blocks -----------------------------------------------------------

    def _title_block(self, slide: Any, slide_spec: SlideSpec) -> None:
        text = slide_spec.title
        if self.defect_on("TY-004", slide_spec.index):
            text = " ".join(
                word if word.lower() in ("a", "the", "to", "is") else word.capitalize()
                for word in text.split()
            )
        title = slide.shapes.title
        if title is None:  # pragma: no cover - all content layouts carry a title
            title = self._text_box(
                slide,
                CONTENT_LEFT_PT,
                TITLE_TOP_PT,
                CONTENT_WIDTH_PT,
                TITLE_HEIGHT_PT,
                "Title",
            )
        self._fill_placeholder(
            title, self._curly(text), self.brand.title_font(), align=PP_ALIGN.LEFT
        )
        if self.defect_on("BR-008", slide_spec.index):
            title.left = Emu(int((CONTENT_LEFT_PT + 24) * 12700))
        self._rule(
            slide,
            CONTENT_LEFT_PT,
            TITLE_RULE_TOP_PT,
            CONTENT_WIDTH_PT,
            self.brand.house_navy,
        )

    def _footnote(self, slide: Any, slide_spec: SlideSpec) -> None:
        text = slide_spec.payload.get("footnote")
        if not text:
            return
        if self.defect_on("CO-009", slide_spec.index):
            # A second as-of date, governing a figure this slide restates
            # correctly. The figures agree, which is exactly why nobody notices:
            # the deck is internally consistent and externally out of date.
            text = f"Source: {self.spec.advisor_mark} analysis as at 30-June-2026."
        top = FOOTNOTE_TOP_PT
        if self.defect_on("LO-005", slide_spec.index):
            top = FOOTNOTE_TOP_PT - 36
        self._rule(
            slide,
            CONTENT_LEFT_PT,
            FOOTNOTE_RULE_TOP_PT,
            288,
            self.brand.rule_grey,
        )
        box = self._text_box(
            slide, CONTENT_LEFT_PT, top, CONTENT_WIDTH_PT, 12, "Footnote"
        )
        self._write_paragraph(
            box.text_frame.paragraphs[0],
            self._curly(str(text)),
            self.brand.footnote_font(),
        )

    def _add_empty_placeholder(self, slide: Any) -> None:
        """Copy an empty body placeholder onto the slide, for HY-006."""
        tree = slide.shapes._spTree
        sp = etree.SubElement(tree, f"{{{P}}}sp")
        nv_sp_pr = etree.SubElement(sp, f"{{{P}}}nvSpPr")
        c_nv_pr = etree.SubElement(nv_sp_pr, f"{{{P}}}cNvPr")
        c_nv_pr.set("id", "900")
        c_nv_pr.set("name", "Empty Content Placeholder")
        etree.SubElement(nv_sp_pr, f"{{{P}}}cNvSpPr")
        nv_pr = etree.SubElement(nv_sp_pr, f"{{{P}}}nvPr")
        ph = etree.SubElement(nv_pr, f"{{{P}}}ph")
        ph.set("type", "body")
        ph.set("idx", "7")
        sp_pr = etree.SubElement(sp, f"{{{P}}}spPr")
        xfrm = etree.SubElement(sp_pr, f"{{{A}}}xfrm")
        off = etree.SubElement(xfrm, f"{{{A}}}off")
        off.set("x", str(720 * 12700))
        off.set("y", str(324 * 12700))
        ext = etree.SubElement(xfrm, f"{{{A}}}ext")
        ext.set("cx", str(204 * 12700))
        ext.set("cy", str(48 * 12700))
        tx_body = etree.SubElement(sp, f"{{{P}}}txBody")
        etree.SubElement(tx_body, f"{{{A}}}bodyPr")
        etree.SubElement(tx_body, f"{{{A}}}p")

    # -- primitives --------------------------------------------------------------

    def _text_box(
        self,
        slide: Any,
        left_pt: float,
        top_pt: float,
        width_pt: float,
        height_pt: float,
        name: str,
        *,
        align: Any = PP_ALIGN.LEFT,
    ) -> Any:
        box = slide.shapes.add_textbox(
            Pt(left_pt), Pt(top_pt), Pt(width_pt), Pt(height_pt)
        )
        box.name = name
        frame = box.text_frame
        frame.word_wrap = False
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        frame.paragraphs[0].alignment = align
        return box

    def _rect(
        self,
        slide: Any,
        left_pt: float,
        top_pt: float,
        width_pt: float,
        height_pt: float,
        fill_hex: str,
        name: str,
    ) -> Any:
        from pptx.enum.shapes import MSO_SHAPE

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Pt(left_pt), Pt(top_pt), Pt(width_pt), Pt(height_pt)
        )
        shape.name = name
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill_hex)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def _rule(
        self,
        slide: Any,
        left_pt: float,
        top_pt: float,
        width_pt: float,
        colour_hex: str,
    ) -> Any:
        """A hairline rule, drawn as a thin filled rectangle.

        A rectangle rather than a connector so its fill participates in palette
        observation the same way every other brand element does.
        """
        shape = self._rect(slide, left_pt, top_pt, width_pt, 1.5, colour_hex, "Rule")
        return shape

    def _placeholder_of(self, slide: Any, ph_type: str) -> Any | None:
        for placeholder in slide.placeholders:
            node = placeholder.element.find(".//p:nvPr/p:ph", NS)
            if node is not None and (node.get("type") or "body") == ph_type:
                return placeholder
        return None

    def _fill_placeholder(
        self, shape: Any, text: str, font: FontSpec, *, align: Any = PP_ALIGN.LEFT
    ) -> None:
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        self._write_paragraph(paragraph, text, font, no_bullet=True)

    def _write_paragraph(
        self,
        paragraph: Any,
        text: str,
        font: FontSpec,
        *,
        bullet_char: str | None = None,
        no_bullet: bool = False,
    ) -> None:
        run = paragraph.add_run()
        run.text = text
        self._apply_font(run.font, font)
        ppr = paragraph._pPr if paragraph._pPr is not None else paragraph._p.get_or_add_pPr()
        for tag in ("buNone", "buChar", "buAutoNum"):
            existing = ppr.find(f"{{{A}}}{tag}")
            if existing is not None:
                ppr.remove(existing)
        if no_bullet or bullet_char is None:
            etree.SubElement(ppr, f"{{{A}}}buNone")
        else:
            ppr.set("marL", str(12 * 12700))
            ppr.set("indent", str(-12 * 12700))
            node = etree.SubElement(ppr, f"{{{A}}}buChar")
            node.set("char", bullet_char)
        _write_run_props(ppr, font)

    def _apply_font(self, font_obj: Any, font: FontSpec) -> None:
        font_obj.name = font.name
        font_obj.size = Pt(font.size_pt)
        font_obj.bold = font.bold
        font_obj.italic = font.italic
        font_obj.color.rgb = _rgb(font.color_hex)

    def _cell_text(
        self, cell: Any, text: str, font: FontSpec, *, align: Any = PP_ALIGN.LEFT
    ) -> None:
        frame = cell.text_frame
        frame.word_wrap = True
        cell.margin_left = Pt(3)
        cell.margin_right = Pt(3)
        cell.margin_top = Pt(1)
        cell.margin_bottom = Pt(1)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        self._write_paragraph(paragraph, text, font, no_bullet=True)

    def _fill_cell(self, cell: Any, colour_hex: str) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(colour_hex)

    def _plain_table_style(self, table: Any) -> None:
        """Strip the built-in table style so the fixture's palette is exactly the
        five colours the spec declares."""
        tbl_pr = table._tbl.find(f"{{{A}}}tblPr")
        if tbl_pr is None:
            return
        tbl_pr.set("firstRow", "0")
        tbl_pr.set("bandRow", "0")
        tbl_pr.set("firstCol", "0")
        tbl_pr.set("bandCol", "0")
        for child in list(tbl_pr):
            if etree.QName(child).localname == "tableStyleId":
                child.text = _PLAIN_TABLE_STYLE
                break
        else:
            node = etree.SubElement(tbl_pr, f"{{{A}}}tableStyleId")
            node.text = _PLAIN_TABLE_STYLE

    def _curly(self, text: str) -> str:
        """Convert straight quotes to curly, honouring the learned convention.

        Applied at generation time so the clean deck is internally consistent and
        TY-001 has a real convention to detect a deviation from.
        """
        text = re.sub(r"(?<=\w)'(?=\w)", "\u2019", text)
        text = re.sub(r"\"([^\"]*)\"", "\u201c\\1\u201d", text)
        return text


def _write_run_props(ppr: etree._Element, font: FontSpec) -> None:
    """Mirror the run font onto the paragraph's ``a:defRPr``.

    Paragraph-level defaults are part of the inheritance chain, and writing them
    gives ``test_inherit`` a real paragraph-level case to resolve rather than a
    hand-forged one.
    """
    existing = ppr.find(f"{{{A}}}defRPr")
    if existing is not None:
        ppr.remove(existing)
    def_rpr = etree.SubElement(ppr, f"{{{A}}}defRPr")
    def_rpr.set("sz", str(round(font.size_pt * 100)))
    def_rpr.set("b", "1" if font.bold else "0")
    def_rpr.set("i", "1" if font.italic else "0")
    solid = etree.SubElement(def_rpr, f"{{{A}}}solidFill")
    srgb = etree.SubElement(solid, f"{{{A}}}srgbClr")
    srgb.set("val", font.color_hex.lstrip("#").upper())
    latin = etree.SubElement(def_rpr, f"{{{A}}}latin")
    latin.set("typeface", font.name)


def _rgb(hex_value: str) -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(  # type: ignore[no-untyped-call]
        hex_value.lstrip("#").upper()
    )


# --------------------------------------------------------------------------------------
# Package rewriting
# --------------------------------------------------------------------------------------


def _rewrite_package(
    path: Path,
    spec: ReferenceSpec,
    *,
    defects: dict[str, int | None],
    variant: str | None,
) -> None:
    """Rewrite the saved archive.

    Two jobs. First, scrub the document metadata python-pptx's template ships
    with, so the clean deck genuinely passes HY-004 rather than passing because
    the rule is lenient. Second, apply the package-level seeded defects, which
    cannot be expressed through python-pptx at all.
    """
    scratch = path.with_suffix(".rebuilding.pptx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(
        scratch, "w", zipfile.ZIP_DEFLATED
    ) as target:
        names = set(source.namelist())
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = _core_xml(spec, populated=variant == "metadata")
            elif item.filename == "docProps/app.xml":
                data = _app_xml(data, populated=variant == "metadata")
            elif item.filename.startswith("ppt/slides/_rels/") and variant == "external_rel":
                data = _add_external_rel(data, item.filename)
            target.writestr(item, data)

        if variant == "comments":
            _write_comment_parts(target, names)

    shutil.move(str(scratch), str(path))
    del defects


def _core_xml(spec: ReferenceSpec, *, populated: bool) -> bytes:
    creator = "A. Analyst" if populated else ""
    modifier = "M. Director" if populated else ""
    category = "Ashcombe internal" if populated else ""
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<cp:coreProperties '
        'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" '
        'xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:dcmitype="http://purl.org/dc/dcmitype/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f"<dc:title>{spec.project_codename}</dc:title>"
        f"<dc:creator>{creator}</dc:creator>"
        f"<cp:lastModifiedBy>{modifier}</cp:lastModifiedBy>"
        f"<cp:category>{category}</cp:category>"
        "<cp:revision>1</cp:revision>"
        "</cp:coreProperties>"
    ).encode()


def _app_xml(original: bytes, *, populated: bool) -> bytes:
    root = etree.fromstring(original)
    ep = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    for tag, value in (
        ("Company", "Ashcombe Partners LLP" if populated else ""),
        ("Manager", "M. Director" if populated else ""),
    ):
        node = root.find(f"{{{ep}}}{tag}")
        if node is None:
            node = etree.SubElement(root, f"{{{ep}}}{tag}")
        node.text = value
    serialised: bytes = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    return serialised


def _add_external_rel(original: bytes, filename: str) -> bytes:
    """Add an external relationship pointing at a local drive path, for HY-007.

    Applied to the first slide's rels only, so the defect is on one slide.
    """
    if not filename.endswith("slide1.xml.rels"):
        return original
    root = etree.fromstring(original)
    ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    node = etree.SubElement(root, f"{{{ns}}}Relationship")
    node.set("Id", "rIdExternal1")
    node.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
    )
    node.set("Target", "file:///C:/Users/analyst/Desktop/meridian_chart.png")
    node.set("TargetMode", "External")
    serialised: bytes = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    return serialised


def _write_comment_parts(target: zipfile.ZipFile, existing: set[str]) -> None:
    """Inject a legacy-shape comment part and its author list, for HY-005."""
    authors = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<p:cmAuthorLst '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:cmAuthor id="1" name="M. Director" initials="MD" lastIdx="1" clrIdx="0"/>'
        "</p:cmAuthorLst>"
    )
    comments = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<p:cmLst '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:cm authorId="1" dt="2026-09-14T09:00:00" idx="1">'
        '<p:pos x="1000" y="1000"/>'
        "<p:text>Confirm this multiple with the client.</p:text>"
        "</p:cm></p:cmLst>"
    )
    if "ppt/commentAuthors.xml" not in existing:
        target.writestr("ppt/commentAuthors.xml", authors)
    target.writestr("ppt/comments/comment1.xml", comments)
